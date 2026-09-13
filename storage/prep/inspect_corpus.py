"""Read verified raw objects inside Cloudflare and return only bounded metadata.

Archives are never extracted to arbitrary paths. SQL is inspected, never executed.
"""

import hashlib
import json
import os
from pathlib import Path
import re
import tarfile
import tempfile
import urllib.request
import zipfile


def request(path, *, data=None, prep=False):
    base = os.environ["SLP_PREP_URL" if prep else "SLP_STORAGE_URL"]
    headers = {
        "Authorization": "Bearer " + os.environ["SLP_STORAGE_TOKEN"],
        "User-Agent": "SLp-Cloud-Storage/1.0",
    }
    if data is not None:
        headers.update(
            {
                "X-Content-SHA256": hashlib.sha256(data).hexdigest(),
                "Content-Type": "application/json",
            }
        )
    return urllib.request.urlopen(
        urllib.request.Request(
            base + path,
            data=data,
            headers=headers,
            method="PUT" if data is not None else "GET",
        ),
        timeout=120,
    )


def save(name, value):
    payload = json.dumps(value, sort_keys=True, default=str, allow_nan=False).encode()
    with request(
        f"/outputs/{os.environ['SLP_JOB']}/{name}.json", data=payload, prep=True
    ) as response:
        response.read(4096)


def download(spec, directory):
    path = Path(directory) / Path(spec["key"]).name
    digest = hashlib.new(spec["checksum"]["algorithm"])
    count = 0
    with request("/objects/" + spec["id"]) as response, path.open("wb") as dest:
        while chunk := response.read(1024 * 1024):
            count += len(chunk)
            if count > spec["bytes"]:
                raise ValueError("Object exceeds manifest")
            digest.update(chunk)
            dest.write(chunk)
    if count != spec["bytes"] or digest.hexdigest() != spec["checksum"]["value"]:
        raise ValueError("Raw object checksum mismatch")
    return path


def text_metadata(stream, name):
    # Enough to see schema/header rows without materializing the archive member.
    head = stream.read(32768).decode("utf-8-sig", errors="replace")
    result = {"name": name, "head_lines": [x[:1000] for x in head.splitlines()[:12]]}
    if name.lower().endswith(".sql"):
        # SQL headers usually contain schema before values. Scan bounded lines only.
        result["create_tables"] = re.findall(r"CREATE TABLE[^;]{0,12000}", head, re.I)
    return result


def xlsx_metadata(file, name):
    import openpyxl

    workbook = openpyxl.load_workbook(file, read_only=True, data_only=True)
    sheets = []
    try:
        for sheet in workbook.worksheets:
            rows = []
            for i, row in enumerate(sheet.iter_rows(values_only=True)):
                if i >= 10:
                    break
                rows.append([str(v)[:300] if v is not None else None for v in row[:50]])
            sheets.append(
                {
                    "name": sheet.title,
                    "rows": sheet.max_row,
                    "columns": sheet.max_column,
                    "head": rows,
                }
            )
    finally:
        workbook.close()
    return {"name": name, "sheets": sheets}


def inspect(path):
    if path.name.endswith(".xlsx"):
        return xlsx_metadata(path, path.name)
    if zipfile.is_zipfile(path):
        with zipfile.ZipFile(path) as archive:
            infos = archive.infolist()
            members = []
            for info in infos[:500]:
                row = {
                    "name": info.filename,
                    "bytes": info.file_size,
                    "compressed_bytes": info.compress_size,
                }
                if not info.is_dir() and not info.filename.startswith("__MACOSX"):
                    with archive.open(info) as stream:
                        row.update(text_metadata(stream, info.filename))
                members.append(row)
            return {
                "member_count": len(infos),
                "members": members,
                "truncated": len(infos) > 500,
            }
    if path.name.endswith((".tar.gz", ".tgz")):
        members = []
        total = 0
        with tarfile.open(path, "r|gz") as archive:
            for info in archive:
                total += 1
                if len(members) >= 500:
                    continue
                row = {
                    "name": info.name,
                    "bytes": info.size,
                    "regular_file": info.isfile(),
                }
                if info.isfile():
                    stream = archive.extractfile(info)
                    row.update(text_metadata(stream, info.name))
                    stream.close()
                members.append(row)
        return {"member_count": total, "members": members, "truncated": total > 500}
    with path.open("rb") as stream:
        return text_metadata(stream, path.name)


def main():
    manifest = json.load(request("/manifest"))
    receipts = []
    for spec in manifest["objects"]:
        with tempfile.TemporaryDirectory() as directory:
            path = Path(directory) / Path(spec["key"]).name
            digest = hashlib.new(spec["checksum"]["algorithm"])
            count = 0
            with request("/objects/" + spec["id"]) as response, path.open("wb") as dest:
                while chunk := response.read(1024 * 1024):
                    count += len(chunk)
                    if count > spec["bytes"]:
                        raise ValueError("Object exceeds manifest")
                    digest.update(chunk)
                    dest.write(chunk)
            if (
                count != spec["bytes"]
                or digest.hexdigest() != spec["checksum"]["value"]
            ):
                raise ValueError("Raw object checksum mismatch")
            report = {"source": spec, "inventory": inspect(path)}
            save(spec["id"], report)
            receipts.append({"id": spec["id"], "verified_bytes": count})
    save(
        "complete",
        {
            "job": os.environ["SLP_JOB"],
            "state": "complete",
            "sources": receipts,
            "training_admission": False,
        },
    )


if __name__ == "__main__":
    try:
        main()
    except Exception as exc:
        # Avoid source URLs, credentials, response bodies and raw rows in logs.
        save("failed", {"state": "failed", "error_type": type(exc).__name__})
        raise SystemExit(1) from None
