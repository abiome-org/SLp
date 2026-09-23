"""Reference baselines for SLB-1.

Unsupervised (no SL labels):
  random            uniform noise
  paralog_identity  max protein sequence identity if the pair are Ensembl 116 paralogs, else 0
  fitness           sickness of the two single mutants: -(f_a + f_b); human = DepMap gene effect
                    in that cell line (else the pan-line mean), yeast = SGA single-mutant fitness
  codependency      correlation of the two genes' DepMap gene-effect profiles (human only)
Supervised (fit on the SLB train split only):
  lgbm              gradient boosting on all features above + species
"""

from __future__ import annotations

import functools
from pathlib import Path

import numpy as np
import polars as pl

from slpbench import ids
from slpbench.evaluate import BENCH, load_split
from slpbench.homology import identity_lookup

RAW = Path("data/raw")
INTERIM = Path("data/interim")


@functools.cache
def depmap() -> tuple[np.ndarray, dict[str, int], dict[str, int]]:
    """(gene effect matrix lines x genes, gene -> column, depmap_id -> row). NaN-filled with gene means."""
    df = pl.read_csv(RAW / "depmap/CRISPRGeneEffect_24Q4.csv")
    rows = {m: i for i, m in enumerate(df[df.columns[0]].to_list())}
    h = ids.human()
    genes = {}
    for j, c in enumerate(df.columns[1:]):
        sym = h(c.split(" (")[0]) or h(c.split("(")[-1].rstrip(")"))
        if sym and sym not in genes:
            genes[sym] = j
    x = df.drop(df.columns[0]).to_numpy().astype(np.float32)
    mu = np.nanmean(x, axis=0)
    x = np.where(np.isnan(x), mu, x)
    return x, genes, rows


@functools.cache
def _depmap_z() -> np.ndarray:
    x, _, _ = depmap()
    return (x - x.mean(0)) / (x.std(0) + 1e-6)


@functools.cache
def yeast_smf() -> dict[str, float]:
    import openpyxl

    wb = openpyxl.load_workbook(INTERIM / "costanzo/S1/strain_ids_and_single_mutant_fitness.xlsx", read_only=True)
    acc: dict[str, list[float]] = {}
    for r in list(wb.worksheets[0].iter_rows(values_only=True))[1:]:
        if r[1] and isinstance(r[5], (int, float)):
            acc.setdefault(r[1], []).append(float(r[5]))
    return {k: float(np.mean(v)) for k, v in acc.items()}


def features(df: pl.DataFrame) -> pl.DataFrame:
    ctx = pl.read_parquet(BENCH / "contexts.parquet").select("context_id", "depmap_id")
    df = df.join(ctx, on="context_id", how="left")
    x, gcol, lrow = depmap()
    z = _depmap_z()
    mean_eff = x.mean(0)
    smf = yeast_smf()
    par = identity_lookup()
    n = df.height
    fa, fb, mfa, mfb, cod, pid = (np.full(n, np.nan) for _ in range(6))
    for i, (sp, a, b, dm) in enumerate(df.select("species", "gene_a", "gene_b", "depmap_id").iter_rows()):
        pid[i] = par.get((sp, a, b), 0.0) if sp != "spne" else np.nan
        if sp == "human":
            ja, jb = gcol.get(a), gcol.get(b)
            r = lrow.get(dm)
            if ja is not None:
                fa[i] = x[r, ja] if r is not None else mean_eff[ja]
                mfa[i] = mean_eff[ja]
            if jb is not None:
                fb[i] = x[r, jb] if r is not None else mean_eff[jb]
                mfb[i] = mean_eff[jb]
            if ja is not None and jb is not None:
                cod[i] = float(z[:, ja] @ z[:, jb]) / z.shape[0]
        elif sp == "scer":
            # SGA fitness is ~1 for wild type; shift to 0 so both species read "0 = neutral, negative = sick"
            fa[i] = smf.get(a, np.nan) - 1
            fb[i] = smf.get(b, np.nan) - 1
    return df.with_columns(
        pl.Series("fit_min", np.fmin(fa, fb)), pl.Series("fit_max", np.fmax(fa, fb)),
        pl.Series("fit_sum", fa + fb), pl.Series("pan_fit_sum", mfa + mfb),
        pl.Series("codependency", cod), pl.Series("paralog_identity", pid),
    )


FEATURES = ["fit_min", "fit_max", "fit_sum", "pan_fit_sum", "codependency", "paralog_identity"]


def run(name: str, split: str) -> pl.DataFrame:
    df = load_split(split).drop("label", strict=False)
    if name == "random":
        s = np.random.default_rng(0).random(df.height)
    else:
        f = features(df)
        if name == "paralog_identity":
            s = f["paralog_identity"].fill_null(0).fill_nan(0).to_numpy()
        elif name == "fitness":
            s = -f["fit_sum"].fill_nan(None).fill_null(0).to_numpy()
        elif name == "codependency":
            s = f["codependency"].fill_nan(None).fill_null(0).to_numpy()
        elif name == "lgbm":
            s = _lgbm(f)
        else:
            raise SystemExit(f"unknown baseline {name}")
    return pl.DataFrame({"example_id": df["example_id"], "score": s})


def _lgbm(f: pl.DataFrame) -> np.ndarray:
    import lightgbm as lgb

    tr = features(pl.read_parquet(BENCH / "train.parquet"))
    # balance species so yeast's 2M rows do not dominate
    tr = pl.concat([d.sample(min(d.height, 150_000), seed=0) for d in tr.partition_by("species")])
    sp = ["human", "scer", "spom", "dmel", "spne"]

    def mat(d):
        cols = [d[c].fill_null(np.nan).to_numpy() for c in FEATURES]
        cols += [(d["species"] == s).to_numpy().astype(float) for s in sp]
        return np.column_stack(cols)

    w = tr.group_by("species").len().rename({"len": "n"})
    tr = tr.join(w, on="species").with_columns((1.0 / pl.col("n")).alias("w"))
    m = lgb.LGBMClassifier(n_estimators=400, learning_rate=0.05, num_leaves=31, min_child_samples=50,
                           subsample=0.8, subsample_freq=1, colsample_bytree=0.8, verbose=-1, random_state=0)
    m.fit(mat(tr), tr["label"].to_numpy(), sample_weight=tr["w"].to_numpy() * tr.height)
    return m.predict_proba(mat(f))[:, 1]
