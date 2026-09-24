"""Additional non-human eukaryote genetic-interaction sources (agent data-eukaryotes, 2026-09-24).

Each public function named after a source key returns the MEASUREMENT schema via `finalize`.
`*_checks()` functions compute the reproducibility evidence reported in notes/data/<key>.md.
Nothing here is wired into build.py; the lead integrates.

Species codes and canonical IDs:
  scer  SGD systematic ORF          spom  PomBase systematic ID
  dmel  FlyBase FBgn                cele  WormBase WBGene ID (slpbench.ids_extra.cele)
  mmus  MGI symbol (slpbench.ids_extra.mmus)

Sources (notes/data/<key>*.md has the evidence and the include/exclude recommendation):
  spom  frost2012                                   include (S < -4)
  scer  scer_emaps (9 E-MAPs as one source), kuzmin2018, kuzmin2020, costanzo2021   include
  dmel  horn2011 (include, small), billmann2016 (exclude: unverifiable)
  cele  byrne2007 (conditional), lehner2006 (exclude: inferred negatives)
  mmus  roguev2013 (include, small), gier2020 (exclude: artefactual calls)
Single-gene effects for new species: cele_single(), mmus_single().
"""

from __future__ import annotations

import io
import re
import tarfile
import zipfile
from pathlib import Path

import numpy as np
import polars as pl

from . import finalize

RAW = Path("data/raw")
INTERIM = Path("data/interim")

# E-MAP S-score cut-offs. SLB's Ryan 2012 rule is S < -3 / |S| < 1. Frost 2012 S-scores are wider
# (sd 1.23 vs 1.04 on shared pairs; 3.7% of pairs < -3 vs 1.3% in Ryan) and its S < -3 calls are
# recovered by Ryan 2012 at only AUROC 0.645, S < -4 at 0.670, so Frost uses S < -4.
EMAP_POS, EMAP_NEG = -3.0, 1.0
FROST_POS = -4.0


# --------------------------------------------------------------------------------------------
# helpers

def _auroc(y, s) -> float | None:
    from sklearn.metrics import roc_auc_score

    y = np.asarray(y)
    if y.sum() < 5 or (1 - y).sum() < 5:
        return None
    return float(roc_auc_score(y, s))


def _pairs_from_matrix(rows: list[tuple[str, list[str]]], cols: list[str], idf) -> pl.DataFrame:
    """Matrix (row label, values) x column labels -> long [q, a, score] (q = column/query, a = row/array)."""
    cid = [idf(c) for c in cols]
    out_q, out_a, out_s = [], [], []
    for rl, vals in rows:
        a = idf(rl)
        if not a:
            continue
        for q, v in zip(cid, vals):
            if q and v.strip():
                try:
                    s = float(v)
                except ValueError:
                    continue
                out_q.append(q)
                out_a.append(a)
                out_s.append(s)
    return pl.DataFrame({"q": out_q, "a": out_a, "score": out_s})


def orientation_rule_auroc(long: pl.DataFrame, pos: float, neg: float) -> dict:
    """Split measurements of the same unordered pair (other orientation or other allele):
    label from one measurement (S < pos vs |S| < neg), score with another independent one."""
    g = long.filter(pl.col("q") != pl.col("a")).with_columns(
        pl.min_horizontal("q", "a").alias("x"), pl.max_horizontal("q", "a").alias("y"))
    g = g.group_by("x", "y").agg(pl.col("score")).filter(pl.col("score").list.len() >= 2)
    s1, s2 = g["score"].list.get(0).to_numpy(), g["score"].list.get(1).to_numpy()
    p, n = s1 < pos, np.abs(s1) < neg
    m = p | n
    return {"pairs_measured_twice": g.height, "pos": int(p.sum()), "auroc": _auroc(p[m], -s2[m]),
            "pearson": float(np.corrcoef(s1, s2)[0, 1])}


def cross_auroc(a: pl.DataFrame, b: pl.DataFrame) -> dict:
    """AUROC of study b's score for study a's labels over shared (gene_a, gene_b) pairs (same species)."""
    j = a.select("gene_a", "gene_b", "label").drop_nulls("label").join(
        b.select("gene_a", "gene_b", pl.col("score").alias("s2"), pl.col("label").alias("l2")), on=["gene_a", "gene_b"])
    y = j["label"].to_numpy()
    both = j.drop_nulls("l2")
    return {"overlap_labelled": j.height, "overlap_pos": int(y.sum()), "auroc": _auroc(y, -j["s2"].to_numpy()),
            "both_labelled": both.height, "both_pos_a": int(both["label"].sum()), "both_pos_b": int(both["l2"].sum()),
            "pos_agree": int(((both["label"] == 1) & (both["l2"] == 1)).sum())}


# --------------------------------------------------------------------------------------------
# S. pombe: Frost et al. 2012 Cell E-MAP

def _pombe_sysid(label: str) -> str | None:
    m = re.search(r"(SP[A-Z]{1,3}[0-9A-Z]*\.[0-9]+[A-Z]?|SPNCRNA\.[0-9]+|SPMIT\.[0-9]+)", label.upper())
    return m.group(1) if m else None


def _frost_cdt(averaged: bool) -> tuple[list[str], list[tuple[str, list[str]]]]:
    """Frost 2012 supplementary CDT (clustered matrix; columns = queries, rows = arrays)."""
    zname, member = (("mmc2_averaged.zip", "Frost_Spombe_GI_Averaged.cdt") if averaged
                     else ("mmc3_unaveraged.zip", "Frost_Spombe_GI_Unaveraged.cdt"))
    with zipfile.ZipFile(RAW / "frost2012_spombe" / zname) as z:
        blob = z.read(z.namelist()[0])
    with tarfile.open(fileobj=io.BytesIO(blob), mode="r:gz") as t:
        m = next(x for x in t.getmembers() if x.name.endswith(member))
        text = t.extractfile(m).read().decode("latin-1")
    lines = [l for l in text.replace("\r\n", "\n").replace("\r", "\n").split("\n") if l.strip()]
    cols = lines[0].split("\t")[4:]
    rows = []
    for l in lines[1:]:
        p = l.split("\t")
        if p[0] in ("EWEIGHT", "AID"):
            continue
        rows.append((p[1], p[4:]))
    return cols, rows


def frost2012() -> pl.DataFrame:
    """Frost et al. 2012 Cell "Functional repurposing revealed by comparing S. pombe and
    S. cerevisiae genetic interactions": E-MAP, ~600 queries x ~1300 arrays (deletions and
    DAmP/degron hypomorphs), S-scores. Averaged file (mmc2); alleles collapsed to the gene.
    Positive S < -4 (stricter than Ryan's -3, see FROST_POS), negative |S| < 1.
    """
    cols, rows = _frost_cdt(averaged=True)
    long = _pairs_from_matrix(rows, cols, _pombe_sysid)
    df = long.rename({"q": "gene_a", "a": "gene_b"}).with_columns(
        pl.lit("spom").alias("species"), pl.lit("frost2012").alias("source"), pl.lit("972h-").alias("context"),
        pl.lit("E-MAP").alias("mechanism"), pl.lit("S").alias("score_name"),
        pl.lit(None, pl.Float64).alias("signif"), pl.lit(None, pl.String).alias("signif_name"),
        pl.when(pl.col("score") < FROST_POS).then(1).when(pl.col("score").abs() < EMAP_NEG).then(0)
        .otherwise(None).cast(pl.Int8).alias("label"),
    )
    return finalize(df, "spom")


def frost2012_checks() -> list[dict]:
    """(1) within-study: independent orientations/alleles in the unaveraged matrix;
    (2) cross-study vs Ryan 2012 (same species/strain background) in both directions."""
    from slpbench import ids
    from slpbench.sources.yeast import ryan2012

    cols, rows = _frost_cdt(averaged=False)
    r = ids.spom()
    long = _pairs_from_matrix(rows, cols, lambda s: r(_pombe_sysid(s)) if _pombe_sysid(s) else None)
    out = []
    for t in (-2.5, -3.0, -4.0):
        out.append({"check": f"frost2012 unaveraged, independent orientation/allele, pos S<{t}",
                    **orientation_rule_auroc(long, t, EMAP_NEG)})
    fr, ry = frost2012(), ryan2012()
    for t in (-3.0, -4.0, -5.0):
        f_t = fr.with_columns(pl.when(pl.col("score") < t).then(1).when(pl.col("score").abs() < EMAP_NEG).then(0)
                              .otherwise(None).cast(pl.Int8).alias("label"))
        out.append({"check": f"frost2012 labels (S<{t}) scored by ryan2012 S", **cross_auroc(f_t, ry)})
    out.append({"check": "ryan2012 labels scored by frost2012 S", **cross_auroc(ry, fr)})
    return out


# --------------------------------------------------------------------------------------------
# D. melanogaster: Horn et al. 2011 Nat Methods (signaling GI map, combinatorial RNAi)

HORN_POS_Q, HORN_NEG_Q = 0.05, 0.25


def _horn_raw(sheet: str = "nrCells") -> pl.DataFrame:
    import pandas as pd

    from slpbench import ids

    x = pd.read_excel(RAW / "horn2011_dmel/nmeth1581_MOESM15.xls", sheet_name=sheet, header=20)
    x = x[~(x.Gene1.astype(str).str.startswith("Ctrl") | x.Gene2.astype(str).str.startswith("Ctrl"))]
    r = ids.dmel()

    def res(fb, cg, sym):
        return r(fb) or r(cg) or r(sym)

    ga = [res(*t) for t in zip(x.FBgn1.astype(str), x.CG1.astype(str), x.Gene1.astype(str))]
    gb = [res(*t) for t in zip(x.FBgn2.astype(str), x.CG2.astype(str), x.Gene2.astype(str))]
    return pl.DataFrame({"gene_a": ga, "gene_b": gb, "score": np.log2(x.pi.to_numpy(float)),
                         "signif": x["q.value.ttest"].to_numpy(float), "main_a": x.main1.to_numpy(float),
                         "main_b": x.main2.to_numpy(float)})


def horn2011() -> pl.DataFrame:
    """Horn et al. 2011 Nat Methods "Mapping of signaling networks through synthetic genetic
    interaction analysis by RNAi": 93 signaling genes, all pairs, D-Mel2 cells, 2 dsRNAs per gene,
    cell-count phenotype, multiplicative model; pi = log2(measured / expected).
    Positive: q (t-test) < 0.05 and pi < 0. Negative: q > 0.25 and |pi| < median |pi|.
    """
    df = _horn_raw()
    med = df["score"].abs().median()
    df = df.with_columns(
        pl.lit("dmel").alias("species"), pl.lit("horn2011").alias("source"), pl.lit("S2").alias("context"),
        pl.lit("RNAi").alias("mechanism"), pl.lit("log2pi_cellcount").alias("score_name"),
        pl.lit("q_ttest").alias("signif_name"),
        pl.when((pl.col("signif") < HORN_POS_Q) & (pl.col("score") < 0)).then(1)
        .when((pl.col("signif") > HORN_NEG_Q) & (pl.col("score").abs() < med)).then(0)
        .otherwise(None).cast(pl.Int8).alias("label"),
    ).drop("main_a", "main_b")
    return finalize(df)


def fitness_only_auroc(df: pl.DataFrame, single: pl.DataFrame) -> dict:
    """AUROC of -(f_a + f_b) for the labels (REPLICATION.md 'fitness-only' diagnostic)."""
    f = dict(zip(single["gene"], single["effect"]))
    d = df.drop_nulls("label")
    s = -(np.array([f.get(g, np.nan) for g in d["gene_a"]], float) + np.array([f.get(g, np.nan) for g in d["gene_b"]], float))
    ok = ~np.isnan(s)
    y = d["label"].to_numpy()[ok]
    return {"covered_frac": float(ok.mean()), "fitness_only_auroc": _auroc(y, s[ok])}


def _horn_wells() -> pl.DataFrame:
    """Well-level log2 pi (cell count) from the RNAinteractMAPK Bioconductor object (Horn 2011 raw data):
    [ga, gb (target TIDs, ga < gb), orient (ga was the template), ra, rb (0/1 dsRNA index), pi1, pi2 (screens), pim]."""
    import warnings

    import rdata

    with warnings.catch_warnings():
        warnings.simplefilter("ignore")
        c = rdata.conversion.convert(rdata.parser.parse_file(RAW / "horn2011_dmel/Dmel2PPMAPK.rda"))["Dmel2PPMAPK"]
    rid2t = dict(zip(c.reagents.RID, c.reagents.TID))
    ridx = {r: i for _, g in c.reagents.groupby("TID") for i, r in enumerate(sorted(g.RID))}
    qd, td = c.queryDesign, c.templateDesign
    tplate = dict(zip(qd.Plate.astype(int), qd.TemplatePlate.astype(int)))
    qmap = {(int(p), int(n)): r for p, n, r in zip(qd.Plate, qd.QueryNr, qd.RID)}
    tmap = {(int(t), w): (r, int(n)) for t, w, r, n in zip(td.TemplatePlate, td.Well, td.RID, td.QueryNr)}
    tr, qr = [], []
    for p, w in zip(np.asarray(c.plate), np.asarray(c.well)):
        r, n = tmap[(tplate[int(p)], w)]
        tr.append(r)
        qr.append(qmap[(int(p), n)])
    ch = list(c.channelNames).index("nrCells")
    pi = c.pi.values[:, :, ch]
    samples = set(c.targets.TID[c.targets.group == "sample"])
    d = pl.DataFrame({"tr": tr, "qr": qr, "pi1": pi[:, 0], "pi2": pi[:, 1], "pim": pi[:, 2]}).with_columns(
        pl.col("tr").replace_strict(rid2t).alias("tg"), pl.col("qr").replace_strict(rid2t).alias("qg"),
        pl.col("tr").replace_strict(ridx).alias("ti"), pl.col("qr").replace_strict(ridx).alias("qi"))
    d = d.filter(pl.col("tg").is_in(samples) & pl.col("qg").is_in(samples) & (pl.col("tg") != pl.col("qg")))
    o = pl.col("tg") < pl.col("qg")
    return d.select(pl.min_horizontal("tg", "qg").alias("ga"), pl.max_horizontal("tg", "qg").alias("gb"), o.alias("orient"),
                    pl.when(o).then("ti").otherwise("qi").alias("ra"), pl.when(o).then("qi").otherwise("ti").alias("rb"),
                    "pi1", "pi2", "pim")


def _tail_rule_auroc(x1: np.ndarray, x2: np.ndarray, q: float) -> float | None:
    """replication.split_rule_auroc: labels from x1 (q-tail vs |x| < median), scored by x2."""
    p = x1 < np.quantile(x1, q)
    n = np.abs(x1) < np.quantile(np.abs(x1), 0.5)
    m = p | n
    return _auroc(p[m], -x2[m])


def horn2011_checks() -> list[dict]:
    """Within-study (replicate screens, orientations, independent dsRNA pairs) and cross-study
    (vs the two excluded fly maps, Fischer 2015 and Heigwer 2023) evidence for Horn 2011."""
    from slpbench.sources.dmel import fischer2015, heigwer2023

    w = _horn_wells()
    g = w.group_by("ga", "gb").agg(
        pl.col("pi1").mean(), pl.col("pi2").mean(),
        pl.col("pim").filter(pl.col("orient")).mean().alias("o1"), pl.col("pim").filter(~pl.col("orient")).mean().alias("o2"),
        pl.col("pim").filter((pl.col("ra") == 0) & (pl.col("rb") == 0)).mean().alias("r1"),
        pl.col("pim").filter((pl.col("ra") == 1) & (pl.col("rb") == 1)).mean().alias("r2")).drop_nulls().drop_nans()
    out = []
    for a, b, what in (("pi1", "pi2", "replicate screen 1 -> 2"), ("pi2", "pi1", "replicate screen 2 -> 1"),
                       ("o1", "o2", "orientation A-template -> A-query"), ("r1", "r2", "dsRNA pair (1,1) -> (2,2)"),
                       ("r2", "r1", "dsRNA pair (2,2) -> (1,1)")):
        x1, x2 = g[a].to_numpy(), g[b].to_numpy()
        for q in (0.02, 0.05):
            out.append({"check": f"horn2011 {what}, {q:.0%} tail", "pairs": g.height, "pearson": float(np.corrcoef(x1, x2)[0, 1]),
                        "auroc": _tail_rule_auroc(x1, x2, q)})
    h, f, hg = horn2011(), fischer2015(), heigwer2023()
    for a, b, na, nb in ((h, hg, "horn2011", "heigwer2023"), (hg, h, "heigwer2023", "horn2011"),
                         (h, f, "horn2011", "fischer2015"), (f, h, "fischer2015", "horn2011"),
                         (f, hg, "fischer2015", "heigwer2023"), (hg, f, "heigwer2023", "fischer2015")):
        out.append({"check": f"{na} labels scored by {nb}", **cross_auroc(a, b)})
    return out


# --------------------------------------------------------------------------------------------
# Single-gene fitness for new species (negative = sicker)

WB_LETHAL = {"WBPhenotype:0000062", "WBPhenotype:0000059"}          # lethal, larval arrest (+ descendants)
WB_SICK = {"WBPhenotype:0000688", "WBPhenotype:0000031"}            # sterile, slow growth (+ descendants)


def _obo_descendants(path: Path, roots: set[str]) -> set[str]:
    children: dict[str, set[str]] = {}
    cur = None
    for line in open(path):
        line = line.strip()
        if line.startswith("id: "):
            cur = line[4:]
        elif line.startswith("is_a: ") and cur:
            children.setdefault(line[6:].split(" ")[0], set()).add(cur)
    out, todo = set(roots), list(roots)
    while todo:
        for c in children.get(todo.pop(), ()):
            if c not in out:
                out.add(c)
                todo.append(c)
    return out


def cele_single() -> pl.DataFrame:
    """C. elegans single-gene loss effect from WormBase WS298 phenotype annotations (RNAi and alleles):
    -1 lethal / larval arrest (any descendant term, positively annotated), -0.5 sterile or slow growth,
    0 otherwise for genes with any phenotype annotation (incl. NOT-annotations, i.e. the gene was assayed).
    Mirrors the S. pombe viability scale used by SLB (fitness.pombe_viability)."""
    obo = RAW / "wormbase/phenotype_ontology.WS298.obo"
    lethal, sick = _obo_descendants(obo, WB_LETHAL), _obo_descendants(obo, WB_SICK)
    d = pl.read_csv(RAW / "wormbase/phenotype_association.WS298.wb", separator="\t", has_header=False, comment_prefix="!",
                    infer_schema_length=0, quote_char=None, columns=[1, 3, 4], new_columns=["gene", "qual", "term"])
    pos = d.filter(pl.col("qual").fill_null("") != "NOT")
    v = pos.with_columns(pl.when(pl.col("term").is_in(list(lethal))).then(-1.0)
                         .when(pl.col("term").is_in(list(sick))).then(-0.5).otherwise(0.0).alias("effect"))
    eff = v.group_by("gene").agg(pl.col("effect").min())
    assayed = d.select("gene").unique().with_columns(pl.lit(0.0).alias("effect"))
    out = pl.concat([eff, assayed.join(eff, on="gene", how="anti")])
    return out.filter(pl.col("gene").str.starts_with("WBGene"))


# --------------------------------------------------------------------------------------------
# D. melanogaster: Billmann et al. 2016 MBoC (cell-cycle GI map, combinatorial RNAi)

BILLMANN_POS, BILLMANN_NEG = -3.0, 1.0


def billmann2016() -> pl.DataFrame:
    """Billmann et al. 2016 Mol Biol Cell "A genetic interaction map of cell cycle regulators":
    350 candidate x 14 query genes, S2 cells, combinatorial RNAi, median pi-score (cell count sheet).
    No p-values are published. Positive pi < -3, negative |pi| < 1."""
    import warnings

    import pandas as pd

    from slpbench import ids

    base = RAW / "billmann2016_dmel/supp_E15-07-0467v1_mc-E15-07-0467-"
    with warnings.catch_warnings():
        warnings.simplefilter("ignore")
        m = pd.read_excel(f"{base}s07.xlsx", sheet_name="gi_cellCount", index_col=0)
        ann = pd.concat([pd.read_excel(f"{base}s05.xlsx", header=8), pd.read_excel(f"{base}s06.xlsx", header=6)])
    r = ids.dmel()
    sym = {str(s): r(str(fb)) or r(str(cg)) or r(str(s)) for fb, cg, s in zip(ann.Gene_FBgn, ann.Gene_CG, ann.Gene_Symbol)}
    long = m.stack().reset_index()
    long.columns = ["cand", "query", "score"]
    df = pl.DataFrame({"gene_a": [sym.get(str(s)) or r(str(s)) for s in long.cand],
                       "gene_b": [sym.get(str(s)) or r(str(s)) for s in long["query"]],
                       "score": long.score.to_numpy(float)})
    df = df.with_columns(
        pl.lit("dmel").alias("species"), pl.lit("billmann2016").alias("source"), pl.lit("S2").alias("context"),
        pl.lit("RNAi").alias("mechanism"), pl.lit("pi_cellcount").alias("score_name"),
        pl.lit(None, pl.Float64).alias("signif"), pl.lit(None, pl.String).alias("signif_name"),
        pl.when(pl.col("score") < BILLMANN_POS).then(1).when(pl.col("score").abs() < BILLMANN_NEG).then(0)
        .otherwise(None).cast(pl.Int8).alias("label"))
    return finalize(df)


def billmann2016_checks() -> list[dict]:
    from slpbench.sources.dmel import fischer2015, heigwer2023

    b = billmann2016()
    out = []
    for n, o in (("fischer2015", fischer2015()), ("heigwer2023", heigwer2023()), ("horn2011", horn2011())):
        out.append({"check": f"billmann2016 labels scored by {n}", **cross_auroc(b, o)})
        out.append({"check": f"{n} labels scored by billmann2016", **cross_auroc(o, b)})
    return out


# --------------------------------------------------------------------------------------------
# C. elegans: Byrne et al. 2007 J Biol (RNAi x query mutant growth), Lehner et al. 2006 Nat Genet

BYRNE_POS = 2.0


def _cele_ids(names) -> list[str | None]:
    from slpbench import ids_extra

    r = ids_extra.cele()
    out = []
    for n in names:
        n = str(n).strip()
        g = r(n)
        if g is None and "_" in n:  # 'locus_sequence' or '_sequence'
            parts = [p for p in n.split("_") if p]
            g = next((r(p) for p in reversed(parts) if r(p)), None)
        if g is None and n[-1:].isalpha() and n[-1:].islower():  # isoform suffix, e.g. AH10.5a
            g = r(n[:-1])
        out.append(g)
    return out


def _byrne_raw() -> pl.DataFrame:
    import warnings

    import pandas as pd

    f = RAW / "byrne2007_cele/jbiol58-S4.xls"
    with warnings.catch_warnings():
        warnings.simplefilter("ignore")
        allx = pd.read_excel(f, sheet_name="All Interaction Strengths", header=5)
        net = pd.read_excel(f, sheet_name="SGI Network Strengths", header=5)
    netk = set(zip(net.Query.astype(str), net.Target.astype(str)))
    return pl.DataFrame({
        "query": allx.Query.astype(str).tolist(), "target": allx.Target.astype(str).tolist(),
        "gene_a": _cele_ids(allx.Query), "gene_b": _cele_ids(allx.Target),
        "strength": allx.Strength.to_numpy(float), "batch": allx.Screen.astype(str).tolist(),
        "in_network": [(q, t) in netk for q, t in zip(allx.Query.astype(str), allx.Target.astype(str))]})


def byrne2007() -> pl.DataFrame:
    """Byrne et al. 2007 J Biol 6:8 "A global analysis of genetic interactions in C. elegans":
    feeding RNAi of ~860 genes (signalling set + chromosome III set) in 11 query mutant strains,
    semi-quantitative growth scoring; interaction strength = enhancement of the RNAi phenotype in the
    mutant over wild type, averaged across rounds (0-6 scale; higher = sicker). Suppl. S4 has all
    ~7,000 tested pairs; the 1,246-edge SGI network lists the pairs that passed the authors'
    internal-consistency criteria. score = -strength (SLB convention: negative = aggravating).
    Positive: in the authors' SGI network and strength >= 2. Negative: not in the network and
    |strength| below the median |strength|."""
    d = _byrne_raw()
    med = d["strength"].abs().median()
    d = d.with_columns(
        (-pl.col("strength")).alias("score"),
        pl.when(pl.col("in_network") & (pl.col("strength") >= BYRNE_POS)).then(1)
        .when(~pl.col("in_network") & (pl.col("strength").abs() < med)).then(0).otherwise(None).cast(pl.Int8).alias("label"),
        pl.lit("cele").alias("species"), pl.lit("byrne2007").alias("source"), pl.lit("N2").alias("context"),
        pl.lit("RNAi x mutant").alias("mechanism"), pl.lit("neg_enhancement_strength").alias("score_name"),
        pl.lit(None, pl.Float64).alias("signif"), pl.lit(None, pl.String).alias("signif_name"))
    return finalize(d)


def lehner2006() -> pl.DataFrame:
    """Lehner et al. 2006 Nat Genet 38:896: feeding RNAi of 1,744 genes (1,860 clones) in query mutant
    strains; binary synthetic-enhancement calls (Emb/Ste/Gro/Lvl/Bmd/Rup). The supplement lists the
    377 hits for the 21 query strains with >= 1 hit and the full clone library. score = -1 for a hit,
    0 otherwise. Positive: listed hit. Negative: bait x library pair not listed (tested, not called;
    no quantitative score, so no neutral band can be applied)."""
    import warnings

    import pandas as pd

    with warnings.catch_warnings():
        warnings.simplefilter("ignore")
        lib = pd.read_excel(RAW / "lehner2006_cele/MOESM1.xls", header=4)
        hits = pd.read_excel(RAW / "lehner2006_cele/MOESM2.xls", header=4)
    clone2g = {str(c): w for c, w in zip(lib.Genepairs_name, lib.WBGene_ID) if str(w).startswith("WBGene")}
    lib_genes = sorted(set(clone2g.values()))
    baits = hits.bait.astype(str).unique().tolist()
    bait = dict(zip(baits, _cele_ids(baits)))
    hit = {(bait[b], clone2g.get(str(c)) or _cele_ids([t])[0])
           for b, c, t in zip(hits.bait.astype(str), hits.target_Ahringer_library_genepairs_ID, hits.target.astype(str))}
    rows = [(q, g, -1.0 if (q, g) in hit else 0.0) for q in sorted({v for v in bait.values() if v}) for g in lib_genes]
    rows += [(q, g, -1.0) for q, g in hit if g not in set(lib_genes)]
    d = pl.DataFrame(rows, schema=["gene_a", "gene_b", "score"], orient="row").with_columns(
        (pl.col("score") < 0).cast(pl.Int8).alias("label"),
        pl.lit("cele").alias("species"), pl.lit("lehner2006").alias("source"), pl.lit("N2").alias("context"),
        pl.lit("RNAi x mutant").alias("mechanism"), pl.lit("hit").alias("score_name"),
        pl.lit(None, pl.Float64).alias("signif"), pl.lit(None, pl.String).alias("signif_name"))
    return finalize(d)


# --------------------------------------------------------------------------------------------
# M. musculus: Roguev et al. 2013 Nat Methods (esiRNA E-MAP, chromatin, fibroblasts),
#              Gier et al. 2020 Nat Commun (multiplexed Cas12a, RN2 AML)

def _mmus_ids(names) -> list[str | None]:
    from slpbench import ids_extra

    r = ids_extra.mmus()
    out = []
    for n in names:
        s = str(n).split(" - ")[0].strip()
        out.append(r(s) or r(s.capitalize()))
    return out


def _roguev_matrix(member: str) -> pl.DataFrame:
    with zipfile.ZipFile(RAW / "roguev2013_mmus/MOESM185.zip") as z:
        name = next(n for n in z.namelist() if n.endswith(member))
        text = z.read(name).decode("latin-1")
    lines = [l.split("\t") for l in text.replace("\r\n", "\n").replace("\r", "\n").split("\n") if l.strip()]
    cols = _mmus_ids(lines[0][1:])
    rows = _mmus_ids([l[0] for l in lines[1:]])
    q, a, s = [], [], []
    for rid, l in zip(rows, lines[1:]):
        for cid, v in zip(cols, l[1:]):
            if rid and cid and v.strip() and v.strip() != "NaN":
                q.append(cid)
                a.append(rid)
                s.append(float(v))
    return pl.DataFrame({"q": q, "a": a, "score": s})


def roguev2013() -> pl.DataFrame:
    """Roguev et al. 2013 Nat Methods 10:432 "Quantitative genetic-interaction mapping in mammalian
    cells": ~11,000 pairwise esiRNA knockdowns of 130 chromatin factors in mouse fibroblasts, cell
    count read-out, E-MAP S-scores (symmetric averaged matrix, S_scoresAvg.txt).
    Positive S < -3, negative |S| < 1 (the SLB E-MAP rule)."""
    long = _roguev_matrix("S_scoresAvg.txt")
    df = long.rename({"q": "gene_a", "a": "gene_b"}).with_columns(
        pl.lit("mmus").alias("species"), pl.lit("roguev2013").alias("source"), pl.lit("fibroblast").alias("context"),
        pl.lit("RNAi").alias("mechanism"), pl.lit("S").alias("score_name"),
        pl.lit(None, pl.Float64).alias("signif"), pl.lit(None, pl.String).alias("signif_name"),
        pl.when(pl.col("score") < EMAP_POS).then(1).when(pl.col("score").abs() < EMAP_NEG).then(0)
        .otherwise(None).cast(pl.Int8).alias("label"))
    return finalize(df)


def roguev2013_checks() -> list[dict]:
    """Within-study: the unaveraged matrix (S_scores.txt) holds both orientations (gene A esiRNA
    as 'query' vs as 'array') as separate measurements."""
    long = _roguev_matrix("S_scores.txt")
    return [{"check": f"roguev2013 orientation split, pos S<{t}", **orientation_rule_auroc(long, t, EMAP_NEG)}
            for t in (-2.0, -3.0)]


def gier2020() -> pl.DataFrame:
    """Gier et al. 2020 Nat Commun 11:3455: multiplexed opAsCas12a dual-crRNA screen of epigenetic
    regulators in the murine MLL-AF9/Nras AML line RN2 (Source Data Fig 3b, gene-pair level:
    expected vs observed depletion, p, q, diff = expected - observed; positive diff = more depleted
    than expected). Experimental x experimental pairs only; pairs with the non-expressed
    negative-control genes are excluded from the output (see notes: they are called as often).
    score = -diff. Positive: q < 0.1 and diff > 0. Negative: q > 0.25 and |diff| < median |diff|."""
    import warnings

    import pandas as pd

    with warnings.catch_warnings():
        warnings.simplefilter("ignore")
        x = pd.read_excel(RAW / "gier2020_mmus/MOESM11.xlsx", sheet_name="Fig 3", header=2).iloc[:, :10]
    x = x.dropna(subset=["guide1.gene"])
    x = x[(x["guide1.type"] == "experimental") & (x["guide2.type"] == "experimental")]
    df = pl.DataFrame({"gene_a": _mmus_ids(x["guide1.gene"]), "gene_b": _mmus_ids(x["guide2.gene"]),
                       "score": -x.diff_value.to_numpy(float), "signif": x.qvalue.to_numpy(float)})
    med = df["score"].abs().median()
    df = df.with_columns(
        pl.lit("mmus").alias("species"), pl.lit("gier2020").alias("source"), pl.lit("RN2").alias("context"),
        pl.lit("CRISPR-Cas12a").alias("mechanism"), pl.lit("neg_diff_expected_observed").alias("score_name"),
        pl.lit("q").alias("signif_name"),
        pl.when((pl.col("signif") < 0.1) & (pl.col("score") < 0)).then(1)
        .when((pl.col("signif") > 0.25) & (pl.col("score").abs() < med)).then(0).otherwise(None).cast(pl.Int8).alias("label"))
    return finalize(df)


def gier2020_checks() -> list[dict]:
    """Negative-control check: pairs of an experimental gene with a non-expressed control gene cannot
    interact genetically; compare how often they are 'called' with experimental pairs."""
    import warnings

    import pandas as pd

    with warnings.catch_warnings():
        warnings.simplefilter("ignore")
        x = pd.read_excel(RAW / "gier2020_mmus/MOESM11.xlsx", sheet_name="Fig 3", header=2).iloc[:, :10]
    x = x.dropna(subset=["guide1.gene"])
    ee = (x["guide1.type"] == "experimental") & (x["guide2.type"] == "experimental")
    en = x["guide1.type"] != x["guide2.type"]
    call = (x.qvalue < 0.1) & (x.diff_value > 0)
    return [{"check": "gier2020 call rate (q<0.1, diff>0), experimental x experimental", "pairs": int(ee.sum()),
             "rate": float(call[ee].mean()), "median_diff": float(x.diff_value[ee].median())},
            {"check": "gier2020 call rate, experimental x negative-control gene", "pairs": int(en.sum()),
             "rate": float(call[en].mean()), "median_diff": float(x.diff_value[en].median())},
            {"check": "AUROC of diff: experimental pairs vs control-gene pairs (0.5 = indistinguishable)",
             "auroc": _auroc(np.r_[np.ones(int(ee.sum())), np.zeros(int(en.sum()))],
                             np.r_[x.diff_value[ee].to_numpy(float), x.diff_value[en].to_numpy(float)])}]


def mmus_single() -> pl.DataFrame:
    """Mouse single-gene loss effect: DepMap 24Q4 Chronos pan-line mean of the human ortholog
    (Alliance combined orthology, best-score AND best-reverse-score pairs only, one human gene per
    mouse gene). There is no genome-wide mouse equivalent of DepMap; this is a documented proxy
    (a permitted single-gene input: DepMap + orthology)."""
    import gzip

    from slpbench import ids, ids_extra
    from slpbench.fitness import depmap_long

    _, pan = depmap_long(())
    eff = dict(zip(pan["gene"], pan["effect"]))
    h, m = ids.human(), ids_extra.mmus()
    pairs: dict[str, set[str]] = {}
    with gzip.open(RAW / "orthology/ORTHOLOGY-ALLIANCE_COMBINED.tsv.gz", "rt") as f:
        for line in f:
            p = line.rstrip("\n").split("\t")
            if len(p) < 13 or p[2] != "NCBITaxon:10090" or p[6] != "NCBITaxon:9606" or p[11] != "Yes" or p[12] != "Yes":
                continue
            mg, hg = m(p[0]) or m(p[1]), h(p[5])
            if mg and hg in eff:
                pairs.setdefault(mg, set()).add(hg)
    rows = [(g, float(eff[next(iter(hs))])) for g, hs in pairs.items() if len(hs) == 1]
    return pl.DataFrame(rows, schema=["gene", "effect"], orient="row")


def byrne2007_checks() -> list[dict]:
    """(1) within-study: pairs measured twice (signalling and chromosome-III batches, or reciprocal
    query/target), label from one measurement with the SLB thresholds, scored by the other;
    (2) cross-study vs Lehner 2006 (same lab, overlapping queries and RNAi library)."""
    r = _byrne_raw().drop_nulls(["gene_a", "gene_b"]).with_columns(
        pl.min_horizontal("gene_a", "gene_b").alias("x"), pl.max_horizontal("gene_a", "gene_b").alias("y"))
    med = r["strength"].abs().median()
    g = r.group_by("x", "y").agg(pl.col("strength")).filter(pl.col("strength").list.len() >= 2)
    s1, s2 = g["strength"].list.get(0).to_numpy(), g["strength"].list.get(1).to_numpy()
    out = []
    for a, b, what in ((s1, s2, "first -> second"), (s2, s1, "second -> first")):
        p, n = a >= BYRNE_POS, np.abs(a) < med
        m = p | n
        out.append({"check": f"byrne2007 duplicate measurements, {what}", "pairs": g.height, "pos": int(p.sum()),
                    "pearson": float(np.corrcoef(s1, s2)[0, 1]), "auroc": _auroc(p[m], b[m])})
    b, l = byrne2007(), lehner2006()
    out.append({"check": "lehner2006 labels scored by byrne2007", **cross_auroc(l, b)})
    out.append({"check": "byrne2007 labels scored by lehner2006 (binary hit)", **cross_auroc(b, l)})
    return out


# --------------------------------------------------------------------------------------------
# S. cerevisiae: Kuzmin et al. 2018 Science (trigenic SGA; digenic control screens)

def _kuzmin_digenic(files: list[tuple[Path, int]], source: str) -> pl.DataFrame:
    """Digenic rows of a Kuzmin-style raw SGA table: query strain 'geneX+YDL227C' (hoΔ control) x array."""
    parts = []
    for path, skip in files:
        lf = pl.scan_csv(path, separator="\t", quote_char=None, infer_schema_length=10000, skip_rows=skip)
        parts.append(lf.filter(pl.col("Combined mutant type") == "digenic").select(
            pl.col("Query strain ID").str.split("_").list.first().str.split("+").alias("qs"),
            pl.col("Array strain ID").str.split("_").list.first().alias("gene_b"),
            pl.col("Raw genetic interaction score (epsilon)").cast(pl.Float64).alias("score"),
            pl.col("P-value").cast(pl.Float64).alias("signif")).collect())
    d = pl.concat(parts)
    d = d.with_columns(pl.col("qs").list.set_difference(pl.lit(["YDL227C"])).alias("qs")) \
        .filter(pl.col("qs").list.len() == 1).with_columns(pl.col("qs").list.first().alias("gene_a")).drop("qs") \
        .drop_nulls(["score", "signif"])
    med = d["score"].abs().median()
    d = d.with_columns(
        pl.lit("scer").alias("species"), pl.lit(source).alias("source"), pl.lit("S288C").alias("context"),
        pl.lit("SGA").alias("mechanism"), pl.lit("epsilon").alias("score_name"), pl.lit("p").alias("signif_name"),
        pl.when((pl.col("score") < -0.2) & (pl.col("signif") < 0.05)).then(1)
        .when((pl.col("signif") > 0.25) & (pl.col("score").abs() < med)).then(0).otherwise(None).cast(pl.Int8).alias("label"))
    return finalize(d, "scer")


def kuzmin2018() -> pl.DataFrame:
    """Kuzmin et al. 2018 Science 360:eaao1729 "Systematic analysis of complex genetic interactions":
    Data File S1, digenic rows only: query single mutants carried as 'geneX + hoΔ' double-marker
    strains (HO, YDL227C, is a neutral locus) crossed to a ~1,200-strain diagnostic array; SGA
    epsilon and p-value (Costanzo 2016 scoring). Independent re-measurement of pairs that are also in
    Costanzo 2016, same lab and method. Label rule = SLB Costanzo 2016 rule: positive eps < -0.2 and
    p < 0.05; negative p > 0.25 and |eps| < median |eps|."""
    return _kuzmin_digenic([(RAW / "kuzmin2018_scer/DataFileS1_Raw_genetic_interaction_dataset.tsv", 0)], "kuzmin2018")


def kuzmin2020() -> pl.DataFrame:
    """Kuzmin et al. 2020 Science 368:eaaz5667 "Exploring whole-genome duplicate gene retention with
    complex genetic interaction analysis": Dryad doi:10.5061/dryad.g79cnp5m9 Tables S1 (main screens)
    and S3 (pilot screens), digenic rows only ('paralogX + hoΔ' query x diagnostic array). Same
    design and rule as kuzmin2018."""
    d = RAW / "kuzmin2020_scer"
    return _kuzmin_digenic([(d / "TableS1.tsv", 1), (d / "TableS3.tsv", 1)], "kuzmin2020")


def _costanzo2016() -> pl.DataFrame:
    p = INTERIM / "measurements/costanzo2016.parquet"
    if p.exists():
        return pl.read_parquet(p)
    from slpbench.sources.yeast import costanzo2016
    return costanzo2016()


def kuzmin2018_checks() -> list[dict]:
    k, c = kuzmin2018(), _costanzo2016()
    return [{"check": "kuzmin2018 labels scored by costanzo2016 epsilon", **cross_auroc(k, c)},
            {"check": "costanzo2016 labels scored by kuzmin2018 epsilon", **cross_auroc(c, k)}]


def kuzmin2020_checks() -> list[dict]:
    k, c, k18 = kuzmin2020(), _costanzo2016(), kuzmin2018()
    return [{"check": "kuzmin2020 labels scored by costanzo2016 epsilon", **cross_auroc(k, c)},
            {"check": "costanzo2016 labels scored by kuzmin2020 epsilon", **cross_auroc(c, k)},
            {"check": "kuzmin2020 labels scored by kuzmin2018 epsilon", **cross_auroc(k, k18)},
            {"check": "kuzmin2018 labels scored by kuzmin2020 epsilon", **cross_auroc(k18, k)}]


# --------------------------------------------------------------------------------------------
# S. cerevisiae E-MAPs (Krogan / Weissman / Walter / Collins labs), S-scores

EMAP_KEEP_ALLELE = re.compile(r"^(|DELETION|DEL|DAMP|KO|TS.*|.*-TS|DAMP.*)$")


def _split_label(s: str) -> tuple[str, str]:
    """'PSK1 - DELETION' -> ('PSK1', 'DELETION'); 'YAL002W' -> ('YAL002W', '')."""
    s = str(s).strip().strip('"')
    if " - " in s:
        n, a = s.split(" - ", 1)
        return n.strip(), a.strip().upper()
    return s, ""


def _emap_finish(long: pl.DataFrame, source: str) -> pl.DataFrame:
    """long [qn, qa, an, aa, score] (names + allele strings) -> MEASUREMENT (E-MAP rule, scer IDs).
    Keeps loss-of-function alleles only (deletion, DAmP, temperature-sensitive)."""
    from slpbench import ids

    ok = pl.col("qa").str.contains(EMAP_KEEP_ALLELE.pattern) & pl.col("aa").str.contains(EMAP_KEEP_ALLELE.pattern)
    d = long.filter(ok).drop_nulls("score").filter(pl.col("score").is_not_nan())
    d = d.with_columns(ids.resolve("scer", d["qn"]).alias("gene_a"), ids.resolve("scer", d["an"]).alias("gene_b"))
    d = d.with_columns(
        pl.lit("scer").alias("species"), pl.lit(source).alias("source"), pl.lit("S288C").alias("context"),
        pl.lit("E-MAP").alias("mechanism"), pl.lit("S").alias("score_name"),
        pl.lit(None, pl.Float64).alias("signif"), pl.lit(None, pl.String).alias("signif_name"),
        pl.when(pl.col("score") < EMAP_POS).then(1).when(pl.col("score").abs() < EMAP_NEG).then(0)
        .otherwise(None).cast(pl.Int8).alias("label"))
    return finalize(d.select("species", "source", "context", "mechanism", "gene_a", "gene_b", "score", "score_name",
                             "signif", "signif_name", "label"))


def _matrix_long(header: list[str], rows: list[list[str]], qsplit=_split_label, asplit=_split_label) -> pl.DataFrame:
    cols = [qsplit(h) for h in header]
    out = {"qn": [], "qa": [], "an": [], "aa": [], "score": []}
    for r in rows:
        an, aa = asplit(r[0])
        for (qn, qa), v in zip(cols, r[1:]):
            v = str(v).strip()
            if not v or v.lower() == "nan":
                continue
            try:
                s = float(v)
            except ValueError:
                continue
            out["qn"].append(qn)
            out["qa"].append(qa)
            out["an"].append(an)
            out["aa"].append(aa)
            out["score"].append(s)
    return pl.DataFrame(out)


def _text_lines(b: bytes) -> list[list[str]]:
    t = b.decode("latin-1").replace("\r\n", "\n").replace("\r", "\n")
    return [l.split("\t") for l in t.split("\n") if l.strip()]


def fiedler2009() -> pl.DataFrame:
    """Fiedler et al. 2009 Cell 136:952 "Functional organization of the S. cerevisiae phosphorylation
    network": signalling E-MAP, 483 x 483 symmetric S-score matrix (mmc3). E-MAP rule (S<-3 / |S|<1)."""
    L = _text_lines((RAW / "fiedler2009_scer/mmc3_Sscore_matrix.txt").read_bytes())
    return _emap_finish(_matrix_long(L[0][1:], L[1:]), "fiedler2009")


def collins2007() -> pl.DataFrame:
    """Collins et al. 2007 Nature 446:806 "Functional dissection of protein complexes involved in
    yeast chromosome biology using a genetic interaction map": 754 x 754 symmetric S-score matrix
    ('Chromosome biology EMAP data.txt'; header rows Gene / Mutation / Marker)."""
    with zipfile.ZipFile(RAW / "collins2007_scer/MOESM264_zip1.zip") as z:
        L = _text_lines(z.read("Supp data - strains and interaction data/Chromosome biology EMAP data.txt"))
    names, muts = L[0][3:], L[1][3:]
    header = [f"{n} - {m}" for n, m in zip(names, muts)]
    rows = [[f"{r[0]} - {r[1]}"] + r[3:] for r in L[3:]]
    return _emap_finish(_matrix_long(header, rows), "collins2007")


def wilmes2008() -> pl.DataFrame:
    """Wilmes et al. 2008 Mol Cell 32:735 "A genetic interaction map of RNA-processing factors reveals
    links between Sem1/Dss1-containing complexes and mRNA export and splicing": 553 x ~550 S-scores
    split over three sheets (mmc3)."""
    import warnings

    import pandas as pd

    with warnings.catch_warnings():
        warnings.simplefilter("ignore")
        sheets = pd.read_excel(RAW / "wilmes2008_scer/mmc3_Sscore_matrix.xls", sheet_name=None, header=None)
    parts = []
    for k, v in sheets.items():
        if not k.startswith("Subset"):
            continue
        v = v.astype(str)
        parts.append(_matrix_long(v.iloc[0, 1:].tolist(), v.iloc[1:].values.tolist()))
    return _emap_finish(pl.concat(parts), "wilmes2008")


def hoppins2011() -> pl.DataFrame:
    """Hoppins et al. 2011 J Cell Biol 195:323 "A mitochondrial-focused genetic interaction map reveals
    a scaffold-like complex required for inner membrane organization" (MITO-MAP): 1,487 arrays x 481
    queries, final averaged S-scores (Table S1, several sheets; header rows ORF / gene / allele)."""
    import warnings

    import pandas as pd

    with warnings.catch_warnings():
        warnings.simplefilter("ignore")
        sheets = pd.read_excel(RAW / "hoppins2011_scer/TableS1.xls", sheet_name=None, header=None)
    parts = []
    for k, v in sheets.items():
        if not k.startswith("Data"):
            continue
        v = v.astype(str)
        header = [f"{o} - {a}" for o, a in zip(v.iloc[0, 3:], v.iloc[2, 3:])]
        rows = [[f"{r[0]} - {r[2]}"] + list(r[3:]) for r in v.iloc[3:].values.tolist()]
        parts.append(_matrix_long(header, rows))
    return _emap_finish(pl.concat(parts), "hoppins2011")


def surma2013() -> pl.DataFrame:
    """Surma et al. 2013 Mol Cell 51:519 "A lipid E-MAP identifies Ubx2 as a critical regulator of lipid
    saturation and lipid bilayer stress": 742 x 742 symmetric S-score matrix (mmc7 xlsx; gene names,
    allele not given in the matrix)."""
    import warnings

    import pandas as pd

    with zipfile.ZipFile(RAW / "surma2013_scer/mmc7_Sscores.zip") as z:
        blob = z.read("S-Scores lipid E-MAP.xlsx")
    with warnings.catch_warnings():
        warnings.simplefilter("ignore")
        v = pd.read_excel(io.BytesIO(blob), sheet_name=0, header=None).astype(str)
    return _emap_finish(_matrix_long(v.iloc[1, 1:].tolist(), v.iloc[2:].values.tolist()), "surma2013")


def zheng2010() -> pl.DataFrame:
    """Zheng et al. 2010 Mol Syst Biol 6:420 "Epistatic relationships reveal the functional organization
    of yeast transcription factors": TF E-MAP, 323 x 323 symmetric S-score matrix (msb201077-s2)."""
    L = _text_lines((RAW / "zheng2010_scer/msb201077-s2_Sscore_matrix.txt").read_bytes())
    return _emap_finish(_matrix_long(L[0][1:], L[1:]), "zheng2010")


def schuldiner2005() -> pl.DataFrame:
    """Schuldiner et al. 2005 Cell 123:507 "Exploration of the function and organization of the yeast
    early secretory pathway through an epistatic miniarray profile": ESP E-MAP, long format
    (ft1/ft2 systematic name, allele, int_score)."""
    d = pl.read_csv(RAW / "schuldiner2005_scer/schuldiner_esp_Sscores.txt", separator="\t", quote_char=None,
                    infer_schema_length=0)
    long = d.select(pl.col("ft1_systematic_name").alias("qn"), pl.col("ft1_allele").str.to_uppercase().alias("qa"),
                    pl.col("ft2_systematic_name").alias("an"), pl.col("ft2_allele").str.to_uppercase().alias("aa"),
                    pl.col("int_score").cast(pl.Float64, strict=False).alias("score"))
    return _emap_finish(long, "schuldiner2005")


def aguilar2010() -> pl.DataFrame:
    """Aguilar et al. 2010 Nat Struct Mol Biol 17:901 "A plasma-membrane E-MAP reveals links of the
    eisosome with sphingolipid metabolism and endosomal trafficking": TreeView .cdt of S-scores
    (gene names; alleles not stated, treated as deletions/DAmP as in the paper)."""
    with zipfile.ZipFile(RAW / "aguilar2010_scer/MOESM19_ESM_treeview_PM_EMAP.zip") as z:
        L = _text_lines(z.read("nsmb.1829-SM1.cdt"))
    rows = [[r[1]] + r[4:] for r in L[1:] if r[0] not in ("AID", "EWEIGHT")]
    return _emap_finish(_matrix_long(L[0][4:], rows), "aguilar2010")


def guenole2013() -> pl.DataFrame:
    """Guenole et al. 2013 Mol Cell 49:346 "Dissection of DNA damage responses using multiconditional
    genetic interaction maps": DNA-repair E-MAP, ~56 queries x ~2,000 arrays; static S-scores in
    untreated (DMSO), MMS, CPT and zeocin (mmc3). Only the untreated map is used here (context S288C)."""
    import openpyxl

    wb = openpyxl.load_workbook(RAW / "guenole2013_scer/mmc3_Sscores_all_pairs.xlsx", read_only=True)
    it = wb.worksheets[0].iter_rows(values_only=True)
    for r in it:
        if r[0] == "Gene 1":
            head = [str(h).strip() if h else "" for h in r]
            break
    iu = head.index("Untreated (DMSO)")
    rows = [(r[0], r[1], r[iu]) for r in it if r[0] and r[1] and isinstance(r[iu], (int, float))]
    long = pl.DataFrame(rows, schema=["qn", "an", "score"], orient="row").with_columns(
        pl.lit("").alias("qa"), pl.lit("").alias("aa"), pl.col("score").cast(pl.Float64))
    return _emap_finish(long.select("qn", "qa", "an", "aa", "score"), "guenole2013")


SCER_EMAPS = ("schuldiner2005", "collins2007", "wilmes2008", "fiedler2009", "zheng2010", "aguilar2010",
              "hoppins2011", "guenole2013", "surma2013")


def scer_emaps() -> pl.DataFrame:
    """The nine S. cerevisiae E-MAPs above as ONE source ('scer_emaps'). The later maps re-used the
    earlier maps' screens for shared genes (pairs shared between two E-MAPs agree at AUROC 0.87-1.00,
    i.e. largely the same measurements), so they must not count as independent studies. Pairs in
    several maps are merged by finalize (mean S; conflicting labels -> null)."""
    parts = [globals()[k]() for k in SCER_EMAPS]
    d = pl.concat(parts).with_columns(pl.lit("scer_emaps").alias("source"))
    return finalize(d)


def hoppins2011_checks() -> list[dict]:
    """Within-study: Hoppins Table S4 = every individual query x array cross (replicate crosses of
    the same query, and both orientations where a gene is query and array)."""
    L = _text_lines((RAW / "hoppins2011_scer/TableS4.txt").read_bytes())
    L = [l for l in L if not l[0].startswith("#")]
    orf, mut = L[0][3:], L[2][3:]
    header = [f"{o} - {m}" for o, m in zip(orf, mut)]
    rows = [[f"{r[0]} - {r[2]}"] + r[3:] for r in L[3:]]
    long = _matrix_long(header, rows)
    ok = pl.col("qa").str.contains(EMAP_KEEP_ALLELE.pattern) & pl.col("aa").str.contains(EMAP_KEEP_ALLELE.pattern)
    long = long.filter(ok).select(pl.col("qn").alias("q"), pl.col("an").alias("a"), "score")
    return [{"check": f"hoppins2011 individual crosses (replicates/orientations), pos S<{t}",
             **orientation_rule_auroc(long, t, EMAP_NEG)} for t in (-2.5, -3.0)]


def scer_emaps_checks() -> list[dict]:
    c = _costanzo2016()
    out = []
    for k in SCER_EMAPS + ("scer_emaps",):
        d = globals()[k]()
        out.append({"check": f"{k} labels scored by costanzo2016 epsilon", **cross_auroc(d, c)})
        out.append({"check": f"costanzo2016 labels scored by {k} S", **cross_auroc(c, d)})
    return out + hoppins2011_checks()


# --------------------------------------------------------------------------------------------
# S. cerevisiae: Costanzo et al. 2021 Science (GI network across 14 environmental conditions)

C21_SHEETS = ("Diagnostic array_complete", "Genome-scale_Benomyl")


def _costanzo2021_raw() -> pl.DataFrame:
    import openpyxl

    wb = openpyxl.load_workbook(RAW / "costanzo2021_scer/DataFileS3_Raw_interaction_dataset.xlsx", read_only=True)
    frames = []
    for sh in C21_SHEETS:
        it = wb[sh].iter_rows(values_only=True)
        head = list(next(it))
        rows = [r for r in it if r[0]]
        frames.append(pl.DataFrame(rows, schema=head, orient="row", infer_schema_length=None)
                      .with_columns(pl.lit(sh).alias("sheet")))
    d = pl.concat(frames, how="diagonal_relaxed")
    f = [c for c in d.columns if "epsilon" in c or "p_value" in c]
    return d.with_columns([pl.col(c).cast(pl.Float64, strict=False) for c in f]).with_columns(
        pl.col("query_orf").str.split("_").list.first().alias("q"),
        pl.col("array_orf").str.split("_").list.first().alias("a"))


def _c21_label(eps: str, p: str) -> pl.Expr:
    return (pl.when((pl.col(eps) < -0.2) & (pl.col(p) < 0.05)).then(1)
            .when((pl.col(p) > 0.25) & (pl.col(eps).abs() < pl.col(eps).abs().median().over("context"))).then(0)
            .otherwise(None).cast(pl.Int8))


def costanzo2021() -> pl.DataFrame:
    """Costanzo et al. 2021 Science 372:eabf8424 "Environmental robustness of the global yeast genetic
    interaction network": 26 query mutants (+ a genome-scale benomyl screen) x ~1,000-strain diagnostic
    array (4,400 arrays for benomyl), SGA in 14 conditions, each with a matched reference screen run
    alongside; two biological replicates each (Data File S3).
    Contexts: 'S288C' = the matched reference condition (all matched-reference measurements of a
    pair, merged by finalize), 'S288C+<condition>' = condition epsilon.
    Label rule = SLB Costanzo 2016 rule on the replicate-mean epsilon and its p-value: positive
    eps < -0.2 and p < 0.05; negative p > 0.25 and |eps| < median |eps| of that context."""
    d = _costanzo2021_raw()
    ref = d.select("q", "a", pl.lit("S288C").alias("context"), pl.col("mean_reference_epsilon").alias("score"),
                   pl.col("reference_p_value").alias("signif"))
    cond = d.select("q", "a", (pl.lit("S288C+") + pl.col("Condition").str.strip_chars()).alias("context"),
                    pl.col("mean_condition_epsilon").alias("score"), pl.col("condition_p_value").alias("signif"))
    x = pl.concat([ref, cond]).drop_nulls(["score", "signif"]).rename({"q": "gene_a", "a": "gene_b"})
    x = x.with_columns(_c21_label("score", "signif").alias("label"),
                       pl.lit("scer").alias("species"), pl.lit("costanzo2021").alias("source"),
                       pl.lit("SGA").alias("mechanism"), pl.lit("epsilon").alias("score_name"),
                       pl.lit("p").alias("signif_name"))
    return finalize(x, "scer")


def costanzo2021_checks() -> list[dict]:
    """(1) replicate 1 -> replicate 2 per condition (split_rule style, SLB thresholds on rep 1);
    (2) reference-condition labels vs Costanzo 2016 (independent screens, same lab)."""
    d = _costanzo2021_raw()
    out = []
    for kind in ("reference", "condition"):
        r1, r2 = d[f"rep1_{kind}_epsilon"].to_numpy(), d[f"rep2_{kind}_epsilon"].to_numpy()
        ok = ~(np.isnan(r1) | np.isnan(r2))
        r1, r2 = r1[ok], r2[ok]
        p, n = r1 < -0.2, np.abs(r1) < np.median(np.abs(r1))
        m = p | n
        out.append({"check": f"costanzo2021 {kind} epsilon, replicate 1 (eps<-0.2) -> replicate 2", "pairs": int(ok.sum()),
                    "pos": int(p.sum()), "pearson": float(np.corrcoef(r1, r2)[0, 1]), "auroc": _auroc(p[m], -r2[m])})
    c21 = costanzo2021()
    ref = c21.filter(pl.col("context") == "S288C")
    c16 = _costanzo2016()
    out.append({"check": "costanzo2021 reference labels scored by costanzo2016", **cross_auroc(ref, c16)})
    out.append({"check": "costanzo2016 labels scored by costanzo2021 reference", **cross_auroc(c16, ref)})
    for ctx in sorted(c21["context"].unique()):
        if ctx == "S288C":
            continue
        x = c21.filter(pl.col("context") == ctx)
        out.append({"check": f"{ctx} labels scored by costanzo2016 (condition vs standard)", **cross_auroc(x, c16)})
    return out
