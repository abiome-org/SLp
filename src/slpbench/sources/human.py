"""Human combinatorial CRISPR screens.

Label policy (per source, documented in BENCHMARK.md):
  positive  = the authors' own SL / strong-negative-GI call (or their published threshold)
  negative  = tested, not called, and in the neutral band of that screen's score
  otherwise = ambiguous (null)
"""

from __future__ import annotations

import sqlite3
from pathlib import Path

import openpyxl
import polars as pl

from . import finalize

RAW = Path("data/raw")
INTERIM = Path("data/interim")

# SLKB studies taken from the authors' original calls. PubMed ID -> (source key, mechanism).
# Paralog screens (Dede, CHyMErA, Parrish, Thompson, Ito) come from the uniform zdLFC re-scoring
# instead. Excluded: 33956155 (Diehl 2021, RPE1) — 63% of tested pairs are called SL, not a
# plausible hit rate for a GI screen; 36060092 (Tang 2022, 22Rv1) — 22% called SL, same concern;
# 29251726 (Najm 2018) has no effect-size column.
SLKB_STUDIES = {
    "26864203": ("wong2016", "CRISPR-KO"),
    "28319085": ("han2017", "CRISPR-KO"),
    "28319113": ("shen2017", "CRISPR-KO"),
    "29452643": ("zhao2018", "CRISPR-KO"),
    "30033366": ("horlbeck2018", "CRISPRi"),
}


def slkb() -> pl.DataFrame:
    con = sqlite3.connect(INTERIM / "slkb/slkb.sqlite")
    q = f"""select study_origin, cell_line_origin, gene_1, gene_2, SL_or_not, SL_score, statistical_score
            from cdko_original_sl_results where study_origin in ({",".join("'%s'" % s for s in SLKB_STUDIES)})"""
    df = pl.read_database(q, con, infer_schema_length=None)
    out = []
    for pmid, sub in df.partition_by("study_origin", as_dict=True).items():
        key, mech = SLKB_STUDIES[pmid[0]]
        for (cl,), s in sub.partition_by("cell_line_origin", as_dict=True).items():
            med = s["SL_score"].abs().median()
            out.append(s.with_columns(
                pl.lit(key).alias("source"), pl.lit(cl).alias("context"), pl.lit(mech).alias("mechanism"),
                pl.when(pl.col("SL_or_not") == "SL").then(1)
                .when((pl.col("SL_or_not") == "Not SL") & (pl.col("SL_score").abs() < med)).then(0)
                .otherwise(None).cast(pl.Int8).alias("label"),
            ))
    df = pl.concat(out).select(
        pl.lit("human").alias("species"), "source", "context", "mechanism",
        pl.col("gene_1").alias("gene_a"), pl.col("gene_2").alias("gene_b"),
        pl.col("SL_score").cast(pl.Float64).alias("score"), pl.lit("author_score").alias("score_name"),
        pl.col("statistical_score").cast(pl.Float64).alias("signif"), pl.lit("author_stat").alias("signif_name"),
        "label",
    )
    return finalize(df, "human")


# Uniform zdLFC re-scoring (Ryan lab 2025 benchmark, figshare 27868350).
# Positive: zdLFC <= -3 (Dede et al. 2020 convention). Negative: |zdLFC| < 1.
ZDLFC_POS, ZDLFC_NEG = -3.0, 1.0
RYANLAB_FILES = {
    # file: (source key, mechanism, {column: cell line})
    "DeDe_zdLFC.csv": ("dede2020", "CRISPR-Cas12a", {"A549": "A549", "HT29": "HT-29", "OVCAR8": "OVCAR-8"}),
    "ChymeraHAP1.csv": ("chymera2020", "CRISPR-Cas12a", {"HAP1.T18": "HAP1"}),
    "ChymeraRPE1.csv": ("chymera2020", "CRISPR-Cas12a", {"RPE1.T24": "RPE1"}),
    "Parrish_Hela.csv": ("parrish2021", "CRISPR-KO", {"HeLa": "HeLa"}),
    "Parrish_PC9.csv": ("parrish2021", "CRISPR-KO", {"PC9": "PC-9"}),
    "Thompson_zdLFC.csv": ("thompson2021", "CRISPR-KO", {"A375_D28": "A375", "RPE_D28": "RPE1", "MEWO_D28": "MeWo"}),
    "ITO.csv": ("ito2021", "CRISPR-KO", {
        "Meljuso": "MEL-JUSO", "GI1_004": "GI-1", "MEL202_003": "MEL202", "PK1": "PK-1", "MEWO": "MeWo",
        "HS944T": "Hs 944.T", "IPC298": "IPC-298", "A549": "A549", "HSC5": "HSC-5", "HS936T": "Hs 936.T",
        "PATU8988S": "PaTu 8988s",
    }),
}


def ryanlab_zdlfc() -> pl.DataFrame:
    out = []
    for fname, (key, mech, cols) in RYANLAB_FILES.items():
        df = pl.read_csv(RAW / "ryanlab2025_bench" / fname)
        pair = df.columns[0]
        df = df.with_columns(pl.col(pair).str.split_exact("_", 1).struct.rename_fields(["gene_a", "gene_b"])).unnest(pair)
        for col, cl in cols.items():
            out.append(df.select(
                "gene_a", "gene_b", pl.col(col).cast(pl.Float64).alias("score"),
                pl.lit(key).alias("source"), pl.lit(cl).alias("context"), pl.lit(mech).alias("mechanism"),
            ))
    df = pl.concat(out).filter(pl.col("score").is_not_null() & pl.col("score").is_not_nan())
    df = df.with_columns(
        pl.lit("human").alias("species"), pl.lit("zdLFC").alias("score_name"),
        pl.lit(None, pl.Float64).alias("signif"), pl.lit(None, pl.String).alias("signif_name"),
        pl.when(pl.col("score") <= ZDLFC_POS).then(1).when(pl.col("score").abs() < ZDLFC_NEG).then(0)
        .otherwise(None).cast(pl.Int8).alias("label"),
    )
    return finalize(df, "human")


def chou2025() -> pl.DataFrame:
    """Chou et al. 2025 (Hart lab, in4mer): 4,540 paralog pairs x 4 lines. Authors' hit: ZdLFC < -2."""
    wb = openpyxl.load_workbook(RAW / "chou2025/chou2025.xlsx", read_only=True)
    rows = list(wb["ZdLFC"].iter_rows(values_only=True))
    header, body = rows[0], rows[1:]
    out = []
    for j, col in enumerate(header[1:], start=1):
        cl = col.replace("_ZdLFC", "")
        for r in body:
            if r[0] and r[j] is not None:
                a, b = r[0].split("_", 1)
                out.append((a, b, cl, float(r[j])))
    df = pl.DataFrame(out, schema=["gene_a", "gene_b", "context", "score"], orient="row")
    df = df.with_columns(
        pl.lit("human").alias("species"), pl.lit("chou2025").alias("source"),
        pl.lit("CRISPR-Cas12a").alias("mechanism"), pl.lit("ZdLFC").alias("score_name"),
        pl.lit(None, pl.Float64).alias("signif"), pl.lit(None, pl.String).alias("signif_name"),
        pl.when(pl.col("score") < -2).then(1).when(pl.col("score").abs() < 1).then(0)
        .otherwise(None).cast(pl.Int8).alias("label"),
    )
    return finalize(df, "human")


def spidr2025() -> pl.DataFrame:
    """Fielden et al. 2025 SPIDR: CRISPRi all-by-all DDR library in RPE1, re-scored from raw counts.

    The published GEMINI sensitive-lethality calls are recovered by a single replicate's additive GI
    at only AUROC 0.62 (docs: REPLICATION.md), so they fail the benchmark's reproducibility bar.
    The counts themselves replicate well, so SPIDR is scored the same way as the paralog screens
    (Dede et al. zdLFC): per replicate, LFC(d14 vs d0) centred on non-targeting pairs; single
    effect f_g = median LFC of gene x non-targeting; GI = LFC - f_a - f_b, median over guide pairs;
    z-scored over all gene pairs. `_mis` (attenuated) guides are separate alleles and excluded.
    Positive: pooled z <= -3 and GI < 0 in both replicates. Negative: |pooled z| < 1.
    """
    gi = spidr_replicate_gi()
    pooled = gi.select(pl.col("gi_rep1"), pl.col("gi_rep2")).mean_horizontal()
    z = (pooled - pooled.mean()) / pooled.std()
    df = gi.with_columns(z.alias("score")).select(
        pl.lit("human").alias("species"), pl.lit("spidr2025").alias("source"), pl.lit("RPE1").alias("context"),
        pl.lit("CRISPRi").alias("mechanism"), "gene_a", "gene_b", "score", pl.lit("zGI_additive").alias("score_name"),
        pl.lit(None, pl.Float64).alias("signif"), pl.lit(None, pl.String).alias("signif_name"),
        pl.when((pl.col("score") <= -3) & (pl.col("gi_rep1") < 0) & (pl.col("gi_rep2") < 0)).then(1)
        .when(pl.col("score").abs() < 1).then(0).otherwise(None).cast(pl.Int8).alias("label"),
    )
    return finalize(df, "human")


def spidr_replicate_gi() -> pl.DataFrame:
    """Per-replicate additive GI for SPIDR RPE1 gene pairs (columns gene_a, gene_b, gi_rep1, gi_rep2)."""
    import numpy as np

    d = pl.read_csv(RAW / "spidr2025/MOESM9_counts.txt", separator="\t")

    def gene(c):
        return (pl.when(pl.col(c).str.starts_with("non_targeting")).then(pl.lit("CONTROL"))
                .when(pl.col(c).str.contains("_mis")).then(pl.lit("MISMATCH"))
                .otherwise(pl.col(c).str.split("_").list.first()))

    d = d.with_columns(gene("sg1").alias("g1"), gene("sg2").alias("g2"))
    t0 = ((d["RPE1_d0_Rep1"] + d["RPE1_d0_Rep2"]) / 2).to_numpy()
    keep = t0 >= np.quantile(t0, 0.02)
    base = np.log2((t0 + 1) / (t0 + 1).sum())
    g1, g2 = d["g1"].to_numpy(), d["g2"].to_numpy()
    out = None
    for i, col in enumerate(["RPE1_d14_Rep1", "RPE1_d14_Rep2"], start=1):
        te = d[col].to_numpy()
        lfc = np.log2((te + 1) / (te + 1).sum()) - base
        lfc = lfc - np.median(lfc[keep & (g1 == "CONTROL") & (g2 == "CONTROL")])
        x = pl.DataFrame({"g1": g1, "g2": g2, "lfc": lfc}).filter(pl.Series(keep))
        x = x.filter((pl.col("g1") != "MISMATCH") & (pl.col("g2") != "MISMATCH"))
        single = pl.concat([
            x.filter(pl.col("g2") == "CONTROL").select(pl.col("g1").alias("g"), "lfc"),
            x.filter(pl.col("g1") == "CONTROL").select(pl.col("g2").alias("g"), "lfc"),
        ]).filter(pl.col("g") != "CONTROL").group_by("g").agg(pl.col("lfc").median().alias("f"))
        f = dict(zip(single["g"], single["f"]))
        du = x.filter((pl.col("g1") != "CONTROL") & (pl.col("g2") != "CONTROL") & (pl.col("g1") != pl.col("g2")))
        gi = du["lfc"].to_numpy() - np.array([f.get(v, np.nan) for v in du["g1"]]) - np.array([f.get(v, np.nan) for v in du["g2"]])
        du = du.with_columns(pl.Series("gi", gi)).drop_nans("gi").with_columns(
            pl.min_horizontal("g1", "g2").alias("gene_a"), pl.max_horizontal("g1", "g2").alias("gene_b"))
        agg = du.group_by("gene_a", "gene_b").agg(pl.col("gi").median().alias(f"gi_rep{i}"))
        out = agg if out is None else out.join(agg, on=["gene_a", "gene_b"])
    return out


def harle2025() -> pl.DataFrame:
    """Harle et al. 2025 Genome Biol: 472 pairs x 27 cancer lines (Sanger).

    Positive: authors' final binary hit matrix (Table S5). Negative: not a hit, fdr > 0.25 and
    |mean_norm_gi| below the median for that line.
    """
    wb = openpyxl.load_workbook(RAW / "harle2025_calls/MOESM1_additional_file1.xlsx", read_only=True)
    s4 = list(wb["Table S4"].iter_rows(min_row=5, values_only=True))
    h = s4[0]
    s4 = pl.DataFrame([dict(zip(h, r)) for r in s4[1:] if r[0]], infer_schema_length=None).select(
        "sorted_gene_pair", "cell_line_label", "depMapID",
        pl.col("mean_norm_gi").cast(pl.Float64), pl.col("fdr").cast(pl.Float64),
    )
    s5 = list(wb["Table S5"].iter_rows(min_row=4, values_only=True))
    lines = s5[0][1:]
    hits = {(r[0], cl) for r in s5[1:] if r[0] for cl, v in zip(lines, r[1:]) if v == 1}
    s4 = s4.with_columns(
        pl.struct("sorted_gene_pair", "cell_line_label")
        .map_elements(lambda s: (s["sorted_gene_pair"], s["cell_line_label"]) in hits, return_dtype=pl.Boolean)
        .alias("hit")
    )
    med = s4.group_by("cell_line_label").agg(pl.col("mean_norm_gi").abs().median().alias("med"))
    s4 = s4.join(med, on="cell_line_label")
    df = s4.with_columns(
        pl.col("sorted_gene_pair").str.split_exact("|", 1).struct.rename_fields(["gene_a", "gene_b"])
    ).unnest("sorted_gene_pair").select(
        pl.lit("human").alias("species"), pl.lit("harle2025").alias("source"),
        (pl.col("cell_line_label") + "|" + pl.col("depMapID").fill_null("")).alias("context"),
        pl.lit("CRISPR-KO").alias("mechanism"), "gene_a", "gene_b",
        pl.col("mean_norm_gi").alias("score"), pl.lit("mean_norm_gi").alias("score_name"),
        pl.col("fdr").alias("signif"), pl.lit("fdr").alias("signif_name"),
        pl.when(pl.col("hit")).then(1)
        .when((pl.col("fdr") > 0.25) & (pl.col("mean_norm_gi").abs() < pl.col("med"))).then(0)
        .otherwise(None).cast(pl.Int8).alias("label"),
    )
    return finalize(df, "human")


def flister2025() -> pl.DataFrame:
    """Flister et al. 2025 Cell Reports (AbbVie): enCas12a paralog library, ~5.1k pairs x 20 lines.

    Table S5 gives observed/expected LFC, diff and diff_z per line; 11 lines also carry the
    authors' 'Lethal' call (diff_z <= -2 plus further filters). Positive: 'Lethal' where the
    call exists, else diff_z <= -2. Negative: |diff_z| < 1. diff_z <= -2 without a 'Lethal'
    call is ambiguous.
    """
    wb = openpyxl.load_workbook(RAW / "flister2025/mmc6.xlsx", read_only=True)
    rows = list(wb.worksheets[0].iter_rows(min_row=2, values_only=True))
    header, body = rows[0], rows[1:]
    out = []
    for start in range(3, len(header), 9):
        name = header[start]
        if not name:
            continue
        cl = name.replace("Gene_A_B", "").strip(" _()")
        has_call = any(r[start + 8] == "Lethal" for r in body if len(r) > start + 8)
        for r in body:
            if len(r) <= start + 7 or not r[0] or not isinstance(r[start + 7], (int, float)):
                continue
            z = float(r[start + 7])
            called = len(r) > start + 8 and r[start + 8] == "Lethal"
            if has_call:
                label = 1 if called else (0 if abs(z) < 1 else None)
            else:
                label = 1 if z <= -2 else (0 if abs(z) < 1 else None)
            out.append((r[0], r[1], cl, z, label))
    df = pl.DataFrame(out, schema={"gene_a": pl.String, "gene_b": pl.String, "context": pl.String,
                                   "score": pl.Float64, "label": pl.Int8}, orient="row")
    df = df.with_columns(
        pl.lit("human").alias("species"), pl.lit("flister2025").alias("source"),
        pl.lit("CRISPR-Cas12a").alias("mechanism"), pl.lit("diff_z").alias("score_name"),
        pl.lit(None, pl.Float64).alias("signif"), pl.lit(None, pl.String).alias("signif_name"),
    )
    return finalize(df, "human")
