"""Najm, DeWeirdt et al. 2023 Nat Commun 14:448 (doi 10.1038/s41467-023-36150-7, CC BY 4.0): orthogonal
SpCas9 (U6) + SaCas9 (H1) "Big Papi" all-by-all screen of 268 chromatin regulators (4 domain-targeted guides per
gene and enzyme), THP-1 (AML) and Reh (B-ALL), day 21 vs plasmid.

Data: Supplementary Data 3 (MOESM5), processed gene-pair table per cell line with three independent library
screens, each already merged over biological duplicates and both Sp/Sa orientations:
  "300k"       268 genes (+EEF2, non-targeting), 35,684 unordered gene pairs per line  -> the SLB measurement
  "40k"        pilot library, 100-gene subset (4,744 overlapping pairs), same lines      -> independent re-screen
  "Validation" 40-gene follow-up library (780 pairs) chosen from 300k hits, in THP-1, Reh and five more AML
               lines (MV4-11, NOMO-1, OCI-AML2, OCI-AML3, P31/FUJ)                     -> independent re-screen
Columns: lfc, base_lfc_a/b (single-gene LFC with controls), pair_z_score (gene-anchored residual z, the
authors' synergy score; negative = more depletion than expected), p-value, fdr_bh.
Per-replicate / per-orientation values are not published: GEO GSE215348 holds the RNA-seq and, for the screens,
only a 40k H1-barcode count matrix (not pair-resolved); the 300k counts are FASTQ-only.
The validation-only AML lines are not loaded as measurements (780 hit-enriched pairs, not a genome-scale or
unbiased pair panel); they are used only as replication evidence.

Label rule (SLB): the authors call synergy at pair z < -4 (their "synergistic lethal" call additionally needs
LFC below the line average; that extra lethality filter is not used here because it adds single-gene fitness
into the label).
  positive: pair_z_score <= -4 (all such pairs have BH FDR < 0.006)
  negative: |pair_z_score| < 1 (implies p > 0.3)
  else null. Self pairs, EEF2 (cutting/essential control) and non-targeting rows are dropped.
score = pair_z_score, signif = fdr_bh.
"""

from __future__ import annotations

from pathlib import Path

import numpy as np
import polars as pl

from slbench.sources import finalize

RAW = Path("data/raw/najm2023")
S3 = RAW / "41467_2023_36150_MOESM5_ESM.xlsx"
CONTEXTS = {"THP1": "THP-1", "REH": "Reh", "MV411": "MV4-11", "NOMO1": "NOMO-1", "OCIAML2": "OCI-AML2",
            "OCIAML3": "OCI-AML3", "P31FUJ": "P31/FUJ"}
MAIN = ("THP-1", "Reh")
POS_Z, NEG_Z = -4.0, 1.0
DROP = {"EEF2", "Non-targeting"}


def sheet(name: str) -> pl.DataFrame:
    """One MOESM5 sheet ('300k', '40k', 'Validation'), controls and self pairs removed, genes stripped."""
    import openpyxl

    ws = openpyxl.load_workbook(S3, read_only=True)[name]
    rows = list(ws.iter_rows(values_only=True))
    d = pl.DataFrame(rows[1:], schema=list(rows[0]), orient="row", infer_schema_length=None)
    d = d.with_columns(pl.col("gene_a", "gene_b").str.strip_chars(),
                       pl.col("condition").replace_strict(CONTEXTS).alias("context"))
    return d.filter((pl.col("gene_a") != pl.col("gene_b")) & ~pl.col("gene_a").is_in(DROP) & ~pl.col("gene_b").is_in(DROP))


def _ordered(d: pl.DataFrame) -> pl.DataFrame:
    """Resolve symbols (as finalize does) and order each pair (gene_a < gene_b)."""
    from slbench import ids

    d = d.with_columns(ids.resolve("human", d["gene_a"]).alias("gene_a"), ids.resolve("human", d["gene_b"]).alias("gene_b"))
    d = d.filter(pl.col("gene_a").is_not_null() & pl.col("gene_b").is_not_null())
    return d.with_columns(pl.min_horizontal("gene_a", "gene_b").alias("gene_a"), pl.max_horizontal("gene_a", "gene_b").alias("gene_b"))


def _label(z: pl.Expr) -> pl.Expr:
    return pl.when(z <= POS_Z).then(1).when(z.abs() < NEG_Z).then(0).otherwise(None).cast(pl.Int8)


def load() -> pl.DataFrame:
    d = sheet("300k")
    df = d.select(
        pl.lit("human").alias("species"), pl.lit("najm2023").alias("source"), "context",
        pl.lit("CRISPR-KO").alias("mechanism"), "gene_a", "gene_b",
        pl.col("pair_z_score").cast(pl.Float64).alias("score"), pl.lit("pair_z_score").alias("score_name"),
        pl.col("fdr_bh").cast(pl.Float64).alias("signif"), pl.lit("fdr_bh").alias("signif_name"),
        _label(pl.col("pair_z_score")).alias("label"),
    )
    return finalize(df, "human")


def singles() -> pl.DataFrame:
    """Screen's own single-gene LFC (gene x controls), per context: [context, gene, f]."""
    d = sheet("300k")
    return pl.concat([d.select("context", pl.col("gene_a").alias("gene"), pl.col("base_lfc_a").alias("f")),
                      d.select("context", pl.col("gene_b").alias("gene"), pl.col("base_lfc_b").alias("f"))]).unique(["context", "gene"])


def replicates() -> pl.DataFrame:
    """Independent library screens as replicate columns (pair_z_score, same line):
    gi_rep1 = 300k library, gi_rep2 = 40k pilot library, gi_rep3 = validation library (null where not screened).
    True biological duplicates / orientations are merged in every published table."""
    def z(name):
        d = _ordered(sheet(name).filter(pl.col("context").is_in(MAIN)))
        return d.select("gene_a", "gene_b", "context", pl.col("pair_z_score").cast(pl.Float64))
    r = z("300k").rename({"pair_z_score": "gi_rep1"})
    r = r.join(z("40k").rename({"pair_z_score": "gi_rep2"}), on=["gene_a", "gene_b", "context"], how="left")
    r = r.join(z("Validation").rename({"pair_z_score": "gi_rep3"}), on=["gene_a", "gene_b", "context"], how="left")
    return r


# ------------------------------------------------------------------------------------------------------
# Checks for data/interim/new_human/najm2023.json


def _auc(y, s) -> float | None:
    from sklearn.metrics import roc_auc_score

    y, s = np.asarray(y), np.asarray(s, float)
    ok = ~np.isnan(s)
    y, s = y[ok], s[ok]
    if y.sum() < 3 or (1 - y).sum() < 3:
        return None
    return float(roc_auc_score(y, s))


def own_replication() -> dict:
    from slbench.replication import replicate_agreement, split_rule_auroc

    r = replicates()
    out = {}
    for col, nm in (("gi_rep2", "40k_library"), ("gi_rep3", "validation_library")):
        res = {}
        for ctx in list(MAIN) + ["pooled"]:
            x = r.filter(pl.col(col).is_not_null())
            if ctx != "pooled":
                x = x.filter(pl.col("context") == ctx)
            x1, x2 = x["gi_rep1"].to_numpy(), x[col].to_numpy()
            y300 = np.where(x1 <= POS_Z, 1, np.where(np.abs(x1) < NEG_Z, 0, -1))
            yo = np.where(x2 <= POS_Z, 1, np.where(np.abs(x2) < NEG_Z, 0, -1))
            k, ko = y300 >= 0, yo >= 0
            res[ctx] = {
                "pairs": x.height,
                **{k_: v for k_, v in replicate_agreement(x.rename({"gi_rep1": "gi_a", col: "gi_b"}), ["a", "b"]).items()
                   if k_ != "pairs"},
                "slb_rule_300k_labels_scored_by_other": _auc(y300[k], -x2[k]), "n_pos_300k": int((y300 == 1).sum()),
                "slb_rule_other_labels_scored_by_300k": _auc(yo[ko], -x1[ko]), "n_pos_other": int((yo == 1).sum()),
                "split_rule_2pct_300k_to_other": split_rule_auroc(x1, x2, pos_q=0.02),
                "split_rule_2pct_other_to_300k": split_rule_auroc(x2, x1, pos_q=0.02),
                "split_rule_z3_300k_to_other": split_rule_auroc(x1, x2, pos_thr=-3.0),
            }
        out[nm] = res
    # validation-only lines: consistency of the other AML lines with the 300k labels (different line, weak)
    v = _ordered(sheet("Validation").filter(~pl.col("context").is_in(MAIN)))
    v = v.group_by("gene_a", "gene_b").agg(pl.col("pair_z_score").mean().alias("zv"))
    m = load().group_by("gene_a", "gene_b").agg(pl.col("label").max()).drop_nulls("label")
    j = m.join(v, on=["gene_a", "gene_b"])
    out["validation_other_aml_lines_mean_z_for_slb_labels"] = {
        "pairs": j.height, "n_pos": int(j["label"].sum()), "auroc": _auc(j["label"].to_numpy(), -j["zv"].to_numpy())}
    return out


def checks() -> dict:
    from slbench import contexts, features

    m = load()
    lab = m.filter(pl.col("label").is_not_null())
    out: dict = {"own_replication": own_replication()}
    genes = set(m["gene_a"]) | set(m["gene_b"])
    par = features.read(Path("data/slb"), "paralogs").filter(pl.col("species") == "human")
    pset = {tuple(sorted(p)) for p in par.select("a", "b").iter_rows()}
    is_par = np.array([(a, b) in pset for a, b in m.select("gene_a", "gene_b").iter_rows()])
    lp = (m["label"] == 1).fill_null(False).to_numpy()
    out["counts"] = {"n_pairs": m.height, "n_unique_pairs": m.select("gene_a", "gene_b").unique().height,
                     "n_pos": int(lp.sum()), "n_neg": int((m["label"] == 0).sum()), "unique_genes": len(genes),
                     "paralog_fraction": float(is_par.mean()), "paralog_pos": int((is_par & lp).sum()),
                     "per_context": {c: {"pairs": g.height, "pos": int((g["label"] == 1).sum()), "neg": int((g["label"] == 0).sum())}
                                     for (c,), g in m.group_by(["context"])}}
    # fitness confounding
    se = pl.read_parquet("data/slb/gene_single_effects.parquet").filter(pl.col("species") == "human")
    f = {g: v for g, v in zip(se["gene"], se["single_effect"]) if v is not None}
    le = pl.read_parquet("data/slb/features/line_effects.parquet")
    s = singles()
    y = lab["label"].to_numpy()
    from slbench import ids  # screen singles use raw symbols; resolve them like finalize does
    res = ids.resolve("human", s["gene"])
    fs = {(c, r_): v for (c, _, v), r_ in zip(s.iter_rows(), res) if r_ and v is not None}
    fl = {}
    for ctx in MAIN:
        dep = contexts.lookup(ctx)["depmap_id"]
        for g, v in le.filter(pl.col("depmap_id") == dep).select("gene", "effect").iter_rows():
            if v is not None:
                fl[(ctx, g)] = v
    rows = list(lab.select("context", "gene_a", "gene_b").iter_rows())
    out["fitness"] = {
        "depmap_mean_sum": _auc(y, -np.array([f.get(a, np.nan) + f.get(b, np.nan) for _, a, b in rows])),
        "depmap_line_sum": _auc(y, -np.array([fl.get((c, a), np.nan) + fl.get((c, b), np.nan) for c, a, b in rows])),
        "screen_singles_sum": _auc(y, -np.array([fs.get((c, a), np.nan) + fs.get((c, b), np.nan) for c, a, b in rows])),
        "paralog_flag": _auc(y, is_par[m["label"].is_not_null().to_numpy()].astype(float)),
    }
    out["cross_study"] = cross_study(m)
    return out


def cross_study(m: pl.DataFrame) -> dict:
    from slbench import contexts

    mine = {c: contexts.lookup(c)["cellosaurus_name"] for c in m["context"].unique()}
    m = m.with_columns(pl.col("context").replace_strict(mine).alias("cname"))
    cs = {}
    for k in ["slkb", "ryanlab_zdlfc", "chou2025", "spidr2025", "harle2025", "flister2025"]:
        o = pl.read_parquet(f"data/interim/measurements/{k}.parquet").filter(pl.col("species") == "human")
        cmap = {c: (contexts.lookup(c.split("|")[0]) or {}).get("cellosaurus_name") for c in o["context"].unique()}
        o = o.with_columns(pl.col("context").replace_strict(cmap, default=None).alias("cname_o"))
        o = o.with_columns((pl.col("score").rank("average").over("source", "context")
                            / pl.col("score").count().over("source", "context")).alias("pct_o"))
        j = m.join(o.select("gene_a", "gene_b", pl.col("source").alias("src_o"), "cname_o", "pct_o",
                            pl.col("label").alias("label_o")), on=["gene_a", "gene_b"])
        if j.height == 0:
            continue
        for same in (True, False):
            jj = j.filter((pl.col("cname") == pl.col("cname_o")) if same else (pl.col("cname") != pl.col("cname_o")))
            for (src,), g in jj.group_by(["src_o"]):
                if not same:  # other lines: one row per pair, mean score/percentile, any-positive labels
                    g = g.group_by("gene_a", "gene_b").agg(pl.col("score").mean(), pl.col("label").max(),
                                                          pl.col("pct_o").mean(), pl.col("label_o").max())
                a = g.filter(pl.col("label").is_not_null())
                b = g.filter(pl.col("label_o").is_not_null())
                cs[f"{src}|{'same_line' if same else 'other_lines'}"] = {
                    "overlap_pairs": g.height,
                    "lines": sorted(set(jj.filter(pl.col("src_o") == src)["cname"])) if same else None,
                    "our_labels_n_pos": int((a["label"] == 1).sum()), "our_labels_n_neg": int((a["label"] == 0).sum()),
                    "auroc_their_score_for_our_labels": _auc(a["label"].to_numpy(), -a["pct_o"].to_numpy()),
                    "their_labels_n_pos": int((b["label_o"] == 1).sum()), "their_labels_n_neg": int((b["label_o"] == 0).sum()),
                    "auroc_our_score_for_their_labels": _auc(b["label_o"].to_numpy(), -b["score"].to_numpy()),
                }
    # pooled over the sources SLB includes (README "Label quality"): same line, and other lines
    inc = {"dede2020", "chou2025", "harle2025", "flister2025", "horlbeck2018", "parrish2021", "zhao2018", "spidr2025"}
    parts = []
    for k in ["slkb", "ryanlab_zdlfc", "chou2025", "spidr2025", "harle2025", "flister2025"]:
        o = pl.read_parquet(f"data/interim/measurements/{k}.parquet").filter(pl.col("source").is_in(inc))
        cmap = {c: (contexts.lookup(c.split("|")[0]) or {}).get("cellosaurus_name") for c in o["context"].unique()}
        parts.append(o.with_columns((pl.col("score").rank("average").over("source", "context")
                                     / pl.col("score").count().over("source", "context")).alias("pct_o"),
                                    pl.col("context").replace_strict(cmap, default=None).alias("cname_o"))
                     .select("gene_a", "gene_b", "cname_o", "pct_o", pl.col("label").alias("label_o")))
    o = pl.concat(parts)
    for same in (True, False):
        if same:
            g = m.join(o, left_on=["gene_a", "gene_b", "cname"], right_on=["gene_a", "gene_b", "cname_o"]) \
                .group_by("gene_a", "gene_b", "cname").agg(pl.col("score").first(), pl.col("label").first(),
                                                          pl.col("pct_o").mean(), pl.col("label_o").max())
        else:
            oo = o.group_by("gene_a", "gene_b").agg(pl.col("pct_o").mean(), pl.col("label_o").max())
            g = m.group_by("gene_a", "gene_b").agg(pl.col("score").mean(), pl.col("label").max()).join(oo, on=["gene_a", "gene_b"])
        a, b = g.filter(pl.col("label").is_not_null()), g.filter(pl.col("label_o").is_not_null())
        cs[f"included_sources_pooled|{'same_line' if same else 'any_line_pair_level'}"] = {
            "overlap_pairs": g.height,
            "our_labels_n_pos": int((a["label"] == 1).sum()), "our_labels_n_neg": int((a["label"] == 0).sum()),
            "auroc_their_score_for_our_labels": _auc(a["label"].to_numpy(), -a["pct_o"].to_numpy()),
            "their_labels_n_pos": int((b["label_o"] == 1).sum()), "their_labels_n_neg": int((b["label_o"] == 0).sum()),
            "auroc_our_score_for_their_labels": _auc(b["label_o"].to_numpy(), -b["score"].to_numpy()),
        }
    return cs
