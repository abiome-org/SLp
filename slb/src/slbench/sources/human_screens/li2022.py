"""Li et al. 2022 Nat Commun 13:2469 (doi 10.1038/s41467-022-30196-9, CC BY 4.0), Sellers lab CombiMiniLib.

Design: 454 paralog pairs (incl. 21 pan-essential and ~111 non-expressed control pairs), 616 genes x 6 guides,
AAVS1 cutting controls. Each gene pair has 18 guide combinations: 3 guides of A (left, U6) x 3 guides of B (right,
H1), plus the reverse orientation with the OTHER 3 guides of each gene (B left x A right). Singles = gene x AAVS1.
Ten library architectures (spCas9 tracrRNA combinations, enCas12a, and spCas9-saCas9 = the published Ito et al.
2021 screens, not in the counts table) in IPC-298; enCas12a, VCR1-WCR3 and WCR2-WCR3 also in PK-1 and MEL-JUSO.
Three biological replicates each. Raw counts: Supplementary Data 9 (MOESM12); pDNA per library is the reference.

Architecture used for the measurement/labels: VCR1-WCR3 (spCas9, alternative tracrRNAs). The paper names it the
best-performing library (strongest pan-essential depletion, most balanced guide positions, most synergistic pairs)
and it is the only top architecture screened in all three lines. The other architectures are exposed by
`architectures()` as a replication check.

Scoring (spidr_replicate_gi recipe, per line and replicate): LFC = log2 norm(end) - log2 norm(pDNA), guide pairs
with pDNA count below the 2% quantile dropped, centred on AAVS1 x AAVS1; f_g = median LFC of gene x AAVS1 guide
pairs; GI = LFC - f_a - f_b; gene-pair GI = median over the 18 guide pairs. Pooled GI uses the mean LFC of the
three replicates; score = z of pooled GI over all gene pairs in that line (mean/sd).

Label rule (same convention as spidr2025 / zdLFC sources):
  positive: z <= -3 and GI < 0 in every replicate
  negative: |z| < 1
  else null.
"""

from __future__ import annotations

from pathlib import Path

import numpy as np
import polars as pl

from slbench.sources import finalize

RAW = Path("data/raw/li2022")
COUNTS_XLSX = RAW / "41467_2022_30196_MOESM12_ESM.xlsx"
SYNERGY_XLSX = RAW / "41467_2022_30196_MOESM14_ESM.xlsx"
CACHE = RAW / "_parsed"

ARCH = "VCR1-WCR3"
LINES = {"IPC298": "IPC-298", "PK1": "PK-1", "MELJUSO": "MEL-JUSO"}
CTRL = "AAVS1"
POS_Z, NEG_Z = -3.0, 1.0


def counts() -> pl.DataFrame:
    """Guide-pair counts with parsed g1 (left), guide index l, g2 (right), guide index r."""
    CACHE.mkdir(exist_ok=True)
    p = CACHE / "counts.parquet"
    if not p.exists():
        import openpyxl

        wb = openpyxl.load_workbook(COUNTS_XLSX, read_only=True)
        rows = list(wb.worksheets[0].iter_rows(min_row=3, values_only=True))
        pl.DataFrame([dict(zip(rows[0], r)) for r in rows[1:] if r[0]], infer_schema_length=None).write_parquet(p)
    d = pl.read_parquet(p)
    lab = d.select(pl.col("pairs").str.extract_groups(r"^(.+)_L(\d+)_(.+)_R(\d+)$")).unnest("pairs")
    lab.columns = ["g1", "l", "g2", "r"]
    return pl.concat([lab, d], how="horizontal")


def _guide_lfc(d: pl.DataFrame, line: str, arch: str) -> tuple[pl.DataFrame, list[str]]:
    reps = sorted(c for c in d.columns if c.startswith(f"{line}_{arch}_Rep"))
    if not reps:
        return None, []
    p0 = d[f"pDNA_{arch}"].cast(pl.Float64).to_numpy()
    keep = p0 >= np.quantile(p0, 0.02)
    base = np.log2((p0 + 1) / (p0 + 1).sum())
    ctrl = keep & (d["g1"] == CTRL).to_numpy() & (d["g2"] == CTRL).to_numpy()
    cols = {}
    for i, c in enumerate(reps, start=1):
        te = d[c].cast(pl.Float64).to_numpy()
        lfc = np.log2((te + 1) / (te + 1).sum()) - base
        cols[f"lfc_rep{i}"] = lfc - np.median(lfc[ctrl])
    x = d.select("g1", "l", "g2", "r").with_columns([pl.Series(k, v) for k, v in cols.items()]).filter(pl.Series(keep))
    x = x.with_columns(pl.mean_horizontal(list(cols)).alias("lfc_mean"))
    return x, list(cols) + ["lfc_mean"]


def gene_pair_gi(line: str, arch: str = ARCH) -> pl.DataFrame | None:
    """Per unordered gene pair: gi_rep1..3, gi_mean (pooled), gi_orient1/2 (A-left vs B-left guide sets, pooled reps)."""
    d = counts()
    x, lcols = _guide_lfc(d, line, arch)
    if x is None:
        return None
    single = pl.concat([
        x.filter((pl.col("g2") == CTRL) & (pl.col("g1") != CTRL)).select(pl.col("g1").alias("g"), *lcols),
        x.filter((pl.col("g1") == CTRL) & (pl.col("g2") != CTRL)).select(pl.col("g2").alias("g"), *lcols),
    ]).group_by("g").agg([pl.col(c).median() for c in lcols])
    # '<GENE>_copy' = a second guide set for the same gene: gene x same-gene cutting controls, not gene pairs
    du = x.filter((pl.col("g1") != CTRL) & (pl.col("g2") != CTRL) & (pl.col("g1") != pl.col("g2"))
                  & ~pl.col("g1").str.ends_with("_copy") & ~pl.col("g2").str.ends_with("_copy"))
    fa = du.select("g1").join(single, left_on="g1", right_on="g", how="left")
    fb = du.select("g2").join(single, left_on="g2", right_on="g", how="left")
    gi = {c.replace("lfc", "gi"): du[c].to_numpy() - fa[c].to_numpy() - fb[c].to_numpy() for c in lcols}
    du = du.with_columns([pl.Series(k, v) for k, v in gi.items()]).with_columns(
        pl.min_horizontal("g1", "g2").alias("gene_a"), pl.max_horizontal("g1", "g2").alias("gene_b"),
        (pl.col("g1") < pl.col("g2")).alias("orient1"))
    gcols = list(gi)
    agg = du.group_by("gene_a", "gene_b").agg(
        *[pl.col(c).drop_nans().median() for c in gcols],
        pl.col("gi_mean").filter(pl.col("orient1")).drop_nans().median().alias("gi_orient1"),
        pl.col("gi_mean").filter(~pl.col("orient1")).drop_nans().median().alias("gi_orient2"),
        pl.len().alias("n_guide_pairs"),
    ).drop_nulls("gi_mean")
    fmap = dict(zip(single["g"], single["lfc_mean"]))
    return agg.with_columns(
        pl.col("gene_a").replace_strict(fmap, default=None).alias("f_a"),
        pl.col("gene_b").replace_strict(fmap, default=None).alias("f_b"),
        pl.lit(LINES[line]).alias("context"),
    ).sort("gene_a", "gene_b")


def scored(line: str, arch: str = ARCH) -> pl.DataFrame:
    g = gene_pair_gi(line, arch)
    rep = [c for c in g.columns if c.startswith("gi_rep")]
    z = (pl.col("gi_mean") - pl.col("gi_mean").mean()) / pl.col("gi_mean").std()
    g = g.with_columns(z.alias("z"))
    all_neg = pl.all_horizontal([pl.col(c) < 0 for c in rep])
    return g.with_columns(
        pl.when((pl.col("z") <= POS_Z) & all_neg).then(1).when(pl.col("z").abs() < NEG_Z).then(0)
        .otherwise(None).cast(pl.Int8).alias("label"))


def load() -> pl.DataFrame:
    df = pl.concat([scored(line) for line in LINES]).select(
        pl.lit("human").alias("species"), pl.lit("li2022").alias("source"), "context",
        pl.lit("CRISPR-KO").alias("mechanism"), "gene_a", "gene_b",
        pl.col("z").alias("score"), pl.lit("zGI_additive").alias("score_name"),
        pl.lit(None, pl.Float64).alias("signif"), pl.lit(None, pl.String).alias("signif_name"), "label",
    )
    return finalize(df, "human")


def replicates() -> pl.DataFrame:
    """VCR1-WCR3, per line: per-biological-replicate GI (gi_rep1..3) and per-orientation GI (gi_orient1/2:
    the two orientations use disjoint guide sets, so they are independent reagent replicates)."""
    return pl.concat([gene_pair_gi(line).select("gene_a", "gene_b", "context", "gi_rep1", "gi_rep2", "gi_rep3",
                                                "gi_orient1", "gi_orient2") for line in LINES])


def architectures() -> pl.DataFrame:
    """Pooled (mean-of-replicates) gene-pair GI for every library architecture with counts, per line.
    Long format: gene_a, gene_b, context, arch, gi. (spCas9-saCas9 = Ito et al. 2021 data, not included.)"""
    d = counts()
    archs = sorted({c.split("_", 1)[1] for c in d.columns if c.startswith("pDNA_")})
    out = []
    for line in LINES:
        for a in archs:
            g = gene_pair_gi(line, a)
            if g is not None:
                out.append(g.select("gene_a", "gene_b", "context", pl.lit(a).alias("arch"), pl.col("gi_mean").alias("gi")))
    return pl.concat(out)


def gemini() -> pl.DataFrame:
    """Authors' GEMINI sensitive synergy scores (Supplementary Data 11), long format."""
    import openpyxl

    wb = openpyxl.load_workbook(SYNERGY_XLSX, read_only=True)
    rows = list(wb.worksheets[0].iter_rows(min_row=3, values_only=True))
    h = rows[0]
    out = []
    for r in rows[1:]:
        if not r[0]:
            continue
        a, b = r[0].split(";")
        for c, v in zip(h[1:], r[1:]):
            if isinstance(v, (int, float)):
                line, arch = c.split("_", 1)
                out.append((min(a, b), max(a, b), LINES.get(line, line), arch, float(v)))
    return pl.DataFrame(out, schema=["gene_a", "gene_b", "context", "arch", "gemini_sens"], orient="row")
