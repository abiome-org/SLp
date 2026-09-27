"""Kim et al. 2025 eLife 93921 (doi 10.7554/eLife.93921, CC BY 4.0): CombiGEM-CRISPR all-by-all screen of
76 tyrosine kinases (+ non-targeting controls, "NegCon") in MDA-MB-231, 3 sgRNAs/gene, both orientations,
three biological replicates (growth phenotype Z r = 0.74 rep2 vs rep3; r = 0.50 between AB/BA orientations).

Data: Supplementary file 2 (elife-93921-supp2-v1.xlsx), one row per UNORDERED gene pair (77*78/2 = 3,003 rows
incl. self pairs and NegCon pairs): observed growth Z ("Normalized Z"), normalized GI (raw GI = deviation from the
quadratic expected-vs-observed Z fit, divided by the SD of the 200 nearest neighbours in expected Z - i.e. a
local z-score), RIGER log10 p for Z and for GI (stored as -log10 p; one-sided, negative direction). The authors'
hits (30 pairs, GI < -2 and RIGER p < 0.01) are bold in the sheet.

Only the gene-level, orientation- and replicate-merged table is published. Raw reads are under BioProject
PRJNA976939, which (Sept 2026) exposes no public runs in SRA/ENA, and no counts table exists; so per-replicate or
per-orientation GI cannot be rebuilt and `replicates()` returns an empty frame (see kim2025.md).

Label rule (SLB):
  positive: authors' call = the 30 bold rows. The paper states "GI < -2 and RIGER p < 0.01", but 40 rows pass
            that; the bold 30 are exactly those that ALSO have observed double-KO growth Z < -5 (the 10 others have
            Z -3.3..-5.0). So positive = GI < -2 AND p(GI) < 0.01 AND Z_obs < -5 (asserted equal to bold); the
            10 unbolded GI/p passers fall to null (they fail the neutral band anyway).
  negative: |normalized GI| < 1 AND RIGER p(GI) > 0.05
  else null. Self pairs (A,A) and NegCon pairs are dropped.
score = normalized GI (negative = synthetic sick/lethal), signif = RIGER p of the GI.
"""

from __future__ import annotations

from pathlib import Path

import numpy as np
import polars as pl

from slbench.sources import finalize

RAW = Path("data/raw/kim2025")
S2 = RAW / "elife-93921-supp2-v1.xlsx"
CONTEXT = "MDA-MB-231"
POS_GI, POS_P, POS_Z = -2.0, 0.01, -5.0
NEG_GI, NEG_P = 1.0, 0.05


def table() -> pl.DataFrame:
    """Raw supp2 rows: x, y (as listed), z_obs, gi, p_z, p_gi, bold (author hit)."""
    import openpyxl

    ws = openpyxl.load_workbook(S2).active
    rows = []
    for r in ws.iter_rows(min_row=5):
        v = [c.value for c in r]
        if not v[0]:
            continue
        x, y = [s.strip() for s in str(v[0]).split(",")]
        rows.append({"x": x, "y": y, "z_obs": float(v[1]), "gi": float(v[2]),
                     "p_z": 10 ** -float(v[3]), "p_gi": 10 ** -float(v[4]), "bold": bool(r[0].font.b)})
    return pl.DataFrame(rows)


def _label(gi: pl.Expr, p: pl.Expr, z: pl.Expr) -> pl.Expr:
    return (pl.when((gi < POS_GI) & (p < POS_P) & (z < POS_Z)).then(1)
            .when((gi.abs() < NEG_GI) & (p > NEG_P)).then(0)
            .otherwise(None).cast(pl.Int8))


def load() -> pl.DataFrame:
    d = table().filter((pl.col("x") != pl.col("y")) & (pl.col("x") != "NegCon") & (pl.col("y") != "NegCon"))
    df = d.select(
        pl.lit("human").alias("species"), pl.lit("kim2025").alias("source"), pl.lit(CONTEXT).alias("context"),
        pl.lit("CRISPR-KO").alias("mechanism"), pl.col("x").alias("gene_a"), pl.col("y").alias("gene_b"),
        pl.col("gi").alias("score"), pl.lit("normalized_GI").alias("score_name"),
        pl.col("p_gi").alias("signif"), pl.lit("RIGER_p_GI").alias("signif_name"),
        _label(pl.col("gi"), pl.col("p_gi"), pl.col("z_obs")).alias("label"),
    )
    assert d.filter(pl.col("bold")).height == df.filter(pl.col("label") == 1).height == 30
    return finalize(df, "human")


def singles() -> pl.DataFrame:
    """The screen's own single-gene growth Z (gene x NegCon row)."""
    d = table().filter((pl.col("x") == "NegCon") ^ (pl.col("y") == "NegCon"))
    return d.select(pl.when(pl.col("x") == "NegCon").then(pl.col("y")).otherwise(pl.col("x")).alias("gene"),
                    pl.col("z_obs").alias("single_z"))


def replicates() -> pl.DataFrame:
    """Not available: only the replicate- and orientation-merged gene-level table is published and the SRA
    BioProject (PRJNA976939) has no public runs. Returns an empty frame with the contract columns."""
    return pl.DataFrame(schema={"gene_a": pl.String, "gene_b": pl.String, "context": pl.String,
                                "gi_rep1": pl.Float64, "gi_rep2": pl.Float64})


# ------------------------------------------------------------------------------------------------------
# Checks for data/interim/new_human/kim2025.json


def _auc(y, s) -> float | None:
    from sklearn.metrics import roc_auc_score

    y, s = np.asarray(y), np.asarray(s, float)
    ok = ~np.isnan(s)
    y, s = y[ok], s[ok]
    if y.sum() < 3 or (1 - y).sum() < 3:
        return None
    return float(roc_auc_score(y, s))


def checks() -> dict:
    from slbench import contexts, features

    m = load()
    lab = m.filter(pl.col("label").is_not_null())
    out: dict = {}
    # counts
    genes = set(m["gene_a"]) | set(m["gene_b"])
    par = features.read(Path("data/slb"), "paralogs").filter(pl.col("species") == "human")
    pset = {tuple(sorted(p)) for p in par.select("a", "b").iter_rows()}
    is_par = np.array([(a, b) in pset for a, b in m.select("gene_a", "gene_b").iter_rows()])
    out["counts"] = {"n_pairs": m.height, "n_pos": int((m["label"] == 1).sum()), "n_neg": int((m["label"] == 0).sum()),
                     "unique_genes": len(genes), "paralog_fraction": float(is_par.mean()),
                     "paralog_pairs": int(is_par.sum()),
                     "paralog_pos": int((is_par & (m["label"] == 1).fill_null(False).to_numpy()).sum())}
    # fitness confounding
    se = pl.read_parquet("data/slb/gene_single_effects.parquet").filter(pl.col("species") == "human")
    f = dict(zip(se["gene"], se["single_effect"]))
    dep = contexts.lookup(CONTEXT)["depmap_id"]
    le = pl.read_parquet("data/slb/features/line_effects.parquet").filter(pl.col("depmap_id") == dep)
    fl = dict(zip(le["gene"], le["effect"]))
    sz = dict(zip(*singles().select("gene", "single_z").to_dict(as_series=False).values()))
    y = lab["label"].to_numpy()

    def fsum(d):
        return -np.array([d.get(a, np.nan) + d.get(b, np.nan) for a, b in lab.select("gene_a", "gene_b").iter_rows()])

    def fmin(d):
        return -np.array([min(d.get(a, np.nan), d.get(b, np.nan)) for a, b in lab.select("gene_a", "gene_b").iter_rows()])

    out["fitness"] = {"depmap_mean_sum": _auc(y, fsum(f)), "depmap_mda_mb_231_sum": _auc(y, fsum(fl)),
                      "screen_singles_sum": _auc(y, fsum(sz)), "screen_singles_min": _auc(y, fmin(sz)),
                      "paralog_flag": _auc(y, is_par[m["label"].is_not_null().to_numpy()].astype(float))}
    # cross-study
    cs = {}
    for k in ["slkb", "ryanlab_zdlfc", "chou2025", "spidr2025", "harle2025", "flister2025"]:
        o = pl.read_parquet(f"data/interim/measurements/{k}.parquet").filter(pl.col("species") == "human")
        cmap = {c: (contexts.lookup(c.split("|")[0]) or {}).get("cellosaurus_name") for c in o["context"].unique()}
        o = o.with_columns(pl.col("context").replace_strict(cmap, default=None).alias("cname"))
        o = o.with_columns((pl.col("score").rank("average").over("source", "context")
                            / pl.col("score").count().over("source", "context")).alias("pct"))
        j = m.join(o.select("gene_a", "gene_b", pl.col("source").alias("src_o"), "cname", pl.col("pct").alias("pct_o"),
                            pl.col("label").alias("label_o"), pl.col("score").alias("score_o")), on=["gene_a", "gene_b"])
        if j.height == 0:
            continue
        for same, jj in ((True, j.filter(pl.col("cname") == CONTEXT)), (False, j.filter(pl.col("cname") != CONTEXT))):
            if jj.height == 0:
                continue
            for (src,), g in jj.group_by(["src_o"]):
                if not same:  # other cell lines: average the other study's percentile over lines
                    g = g.group_by("gene_a", "gene_b").agg(pl.col("score").first(), pl.col("label").first(),
                                                          pl.col("pct_o").mean(), pl.col("label_o").max())
                a = g.filter(pl.col("label").is_not_null())
                b = g.filter(pl.col("label_o").is_not_null())
                cs[f"{src}|{'same_line' if same else 'other_lines'}"] = {
                    "overlap_pairs": g.height,
                    "our_labels_n_pos": int((a["label"] == 1).sum()), "our_labels_n_neg": int((a["label"] == 0).sum()),
                    "auroc_their_score_for_our_labels": _auc(a["label"].to_numpy(), -a["pct_o"].to_numpy()),
                    "their_labels_n_pos": int((b["label_o"] == 1).sum()), "their_labels_n_neg": int((b["label_o"] == 0).sum()),
                    "auroc_our_score_for_their_labels": _auc(b["label_o"].to_numpy(), -b["score"].to_numpy()),
                    "spearman_score": float(g.select(pl.corr("score", "pct_o", method="spearman")).item())
                    if g.height > 5 else None,
                }
    out["cross_study"] = cs
    return out
